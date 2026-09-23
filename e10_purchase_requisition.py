# -*- coding: utf-8 -*-
"""E10 purchase-requisition UI workflow and its operational CLI.

The business scope of this module is limited to ``维护请购单``: input
validation, form entry, one-shot save, identity reconciliation, diagnostics,
and test-ledger tooling.  Future ERP workflows may reuse its proven patterns
and :mod:`rpa_core`, but must not be added to this module.
"""

import argparse
import ctypes
import hashlib
import json
import os
import re
import subprocess
import sys
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

from rpa_core import (REQUIRED_HEADERS, GridSnapshot, QueryOutcome, QueryStatus,
                      ReconcileEvidence,
                      ReconcileStatus, RequisitionRecord, RiskLevel, RunJournal,
                      StepExecutor, StepSpec, StepStatus, WriteIntent,
                      ProgressHeartbeat, derive_ledger,
                      execute_verify_protocol, get_workflow, validate_ledger,
                      enforce_identity_strength,
                      workflow_plan, WindowRecoveryError,
                      plan_window_recovery, FailureCode, GateAction,
                      RequestRegistry, RequestStatus, build_campaign_report,
                      compute_code_hash, load_manual_registrations,
                      merge_manual_registrations, resolve_request_no,
                      scaffold_manual_registrations,
                      IdentityStrength, correlation_value,
                      identity_key, require_dataset_epoch)
from rpa_core import OrphanAction, audit_orphaned_runs


ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "info" / "请购RPA写入测试.xlsx"
DEFAULT_RECORDINGS = ROOT / "runs" / "recordings"
PROBE_MAX_DEPTH = 16
PROBE_NODE_LIMIT = 800
WORKFLOW_VERSION = "2026-09-22.1"
ITEM_LOOKUP_TITLE = "查询品号"
ITEM_LOOKUP_HOTKEY = "{F2}"
UNEXPECTED_ITEM_LOOKUP_TITLES = ("查询资产类品号",)
REQUEST_REGISTRY = ROOT / "runs" / "requisition_requests.json"
MANUAL_REGISTRATIONS = ROOT / "runs" / "requisition_manual_registrations.json"
_RUNTIME_METADATA = {
    "campaign_id": "",
    "code_hash": None,
    "attest_no_intervention": False,
    "scenario_id": "",
    "expected_result": "",
    "desktop_guard": None,
    "lock_namespace": None,
    "mutex_denied_reason": None,
}
_ACTIVE_RUN_JOURNAL: RunJournal | None = None


class _ForegroundTracker:
    """Win32-only observer; it never calls UIA or changes foreground state."""

    def __init__(self, e10_pids: list[int]):
        self.e10_pids = set(e10_pids)
        self._stop = threading.Event()
        self._lock = threading.Lock()
        self._count = 0
        self._saw_e10 = False
        self._foreign_active = False
        self._thread = threading.Thread(
            target=self._run, name="e10-foreground-observer", daemon=True)

    def start(self):
        self._thread.start()

    def _run(self):
        from ctypes import wintypes

        user32 = ctypes.windll.user32
        while not self._stop.wait(0.25):
            hwnd = int(user32.GetForegroundWindow() or 0)
            pid = wintypes.DWORD()
            user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
            is_e10 = int(pid.value) in self.e10_pids
            with self._lock:
                if is_e10:
                    self._saw_e10 = True
                    self._foreign_active = False
                elif self._saw_e10 and not self._foreign_active:
                    self._count += 1
                    self._foreign_active = True

    def finish(self) -> dict:
        self._stop.set()
        self._thread.join(timeout=1.0)
        with self._lock:
            return {"foreign_foreground_observed": self._count}


def _process_creation_time(pid: int) -> int | None:
    """Return the Windows FILETIME creation value without a new dependency."""
    from ctypes import wintypes

    kernel32 = ctypes.windll.kernel32
    kernel32.OpenProcess.restype = wintypes.HANDLE
    kernel32.OpenProcess.argtypes = (
        wintypes.DWORD, wintypes.BOOL, wintypes.DWORD)
    kernel32.GetProcessTimes.argtypes = (
        wintypes.HANDLE, ctypes.POINTER(wintypes.FILETIME),
        ctypes.POINTER(wintypes.FILETIME), ctypes.POINTER(wintypes.FILETIME),
        ctypes.POINTER(wintypes.FILETIME))
    kernel32.CloseHandle.argtypes = (wintypes.HANDLE,)
    handle = kernel32.OpenProcess(0x1000, False, int(pid))
    if not handle:
        return None
    creation = wintypes.FILETIME()
    exit_time = wintypes.FILETIME()
    kernel_time = wintypes.FILETIME()
    user_time = wintypes.FILETIME()
    try:
        if not kernel32.GetProcessTimes(
                handle, ctypes.byref(creation), ctypes.byref(exit_time),
                ctypes.byref(kernel_time), ctypes.byref(user_time)):
            return None
        return (int(creation.dwHighDateTime) << 32) | int(creation.dwLowDateTime)
    finally:
        kernel32.CloseHandle(handle)


def _initial_state_snapshot() -> dict:
    """Read-only Win32 snapshot taken before any workflow action."""
    from ctypes import wintypes
    from e10_query import environment_info, get_digiwin_pids

    user32 = ctypes.windll.user32
    foreground = int(user32.GetForegroundWindow() or 0)
    pids = sorted(int(pid) for pid in get_digiwin_pids())
    windows = []

    enum_proc_type = ctypes.WINFUNCTYPE(
        wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)

    @enum_proc_type
    def collect(hwnd, _lparam):
        hwnd = int(hwnd)
        pid = wintypes.DWORD()
        user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
        if int(pid.value) not in pids:
            return True
        length = user32.GetWindowTextLengthW(hwnd)
        title = ctypes.create_unicode_buffer(length + 1)
        user32.GetWindowTextW(hwnd, title, length + 1)
        rect = wintypes.RECT()
        user32.GetWindowRect(hwnd, ctypes.byref(rect))
        windows.append({
            "hwnd": hwnd,
            "pid": int(pid.value),
            "title": title.value,
            "visible": bool(user32.IsWindowVisible(hwnd)),
            "iconic": bool(user32.IsIconic(hwnd)),
            "foreground": hwnd == foreground,
            "rect": [rect.left, rect.top, rect.right, rect.bottom],
        })
        return True

    user32.EnumWindows(collect, 0)
    material_windows = [row for row in windows if row["visible"]]
    main = [row for row in material_windows
            if row["title"].startswith("鼎捷ERP E10")]
    browse = [row for row in material_windows
              if row["title"].startswith("浏览 - 维护请购单")]
    edit = [row for row in material_windows if row["title"] == "维护请购单"]
    known_prefixes = ("鼎捷ERP E10", "浏览 - 维护请购单", "高级查询")
    known_exact = {"维护请购单", ITEM_LOOKUP_TITLE, *UNEXPECTED_ITEM_LOOKUP_TITLES}
    unknown = [row for row in material_windows if (
        row["title"] and row["title"] not in known_exact
        and not row["title"].startswith(known_prefixes))]
    session = [{"pid": pid, "process_start_filetime": _process_creation_time(pid)}
               for pid in pids]
    return {
        "e10_pids": pids,
        "main_window_count": len(main),
        "browse_window_count": len(browse),
        "edit_window_count": len(edit),
        "draft_present": bool(edit),
        "unknown_windows": unknown,
        "windows": windows,
        "display": environment_info(),
        "client_session": session,
    }


def _make_journal(directory: Path, *, workflow_id: str, selector_version: str,
                  risk: RiskLevel, environment: dict) -> RunJournal:
    global _ACTIVE_RUN_JOURNAL
    merged = {**environment, **_RUNTIME_METADATA}
    snapshot = _initial_state_snapshot()
    merged["client_session"] = snapshot["client_session"]
    tracker = _ForegroundTracker(snapshot["e10_pids"])
    tracker.start()
    journal = RunJournal(
        directory, workflow_id=workflow_id, workflow_version=WORKFLOW_VERSION,
        selector_version=selector_version, risk=risk, environment=merged,
        initial_state=snapshot, terminal_observer=tracker.finish)
    unknown = snapshot.get("unknown_windows") or []
    journal.emit(
        "scene_snapshot", "INVALID" if unknown else "OK",
        ("首次 UI 动作前发现未知 E10 窗口" if unknown
         else "首次 UI 动作前场景快照符合准入条件"),
        scene_code=("PREEXISTING_UNKNOWN_WINDOW" if unknown else "SCENE_READY"),
        sample_eligible=not bool(unknown),
        unknown_windows=unknown,
        main_window_count=snapshot.get("main_window_count"),
        browse_window_count=snapshot.get("browse_window_count"),
        edit_window_count=snapshot.get("edit_window_count"))
    _ACTIVE_RUN_JOURNAL = journal
    return journal


def _step_timing_sink(journal: RunJournal):
    def sink(step: str, status: str, elapsed_ms: float) -> None:
        journal.emit(
            "step_timing", status, step, step=step,
            step_elapsed_ms=round(elapsed_ms, 3))
    return sink


def _record_request_gate_terminal(report_dir: Path, *, status: str,
                                  message: str, request_no: str | None,
                                  payload_fingerprint: str,
                                  dataset_epoch: str | None = None,
                                  doc_no: str | None = None,
                                  existing_evidence: str | None = None) -> Path:
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.request_gate",
        selector_version=selector_version, risk=RiskLevel.COMMIT,
        environment={
            "account_set": "FRKTEST", "mode": "request_gate_only",
            "request_no": request_no,
            "dataset_epoch": dataset_epoch,
            "payload_fingerprint": payload_fingerprint,
        },
    )
    failed = status == "FAILED"
    journal.emit(
        "request_gate", "FAILED" if failed else "OK", message,
        request_no=request_no, doc_no=doc_no,
        dataset_epoch=dataset_epoch,
        existing_evidence=existing_evidence)
    journal.emit(
        "run_finished", status, message,
        business_data_modified=False, safe_to_retry=True,
        request_no=request_no, doc_no=doc_no, deduplicated=not failed,
        dataset_epoch=dataset_epoch,
        existing_evidence=existing_evidence,
        failure_code=FailureCode.INPUT_INVALID.value if failed else None)
    return journal.path


def _record_desktop_guard_result(report_dir: Path, *, status: str, code: str,
                                 message: str, details: dict) -> Path:
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.desktop_guard",
        selector_version=selector_version, risk=RiskLevel.READ_ONLY,
        environment={"mode": "desktop_guard_check"},
    )
    failed = status == "FAILED"
    journal.emit(
        "desktop_guard", "FAILED" if failed else "OK", message,
        guard_code=code if failed else None,
        desktop_context=details,
        business_data_modified=False, safe_to_retry=True)
    journal.emit(
        "run_finished", status, message,
        failure_code=code if failed else None,
        desktop_context=details,
        business_data_modified=False, safe_to_retry=True)
    return journal.path


def _record_desktop_guard_terminal(report_dir: Path, error) -> Path:
    return _record_desktop_guard_result(
        report_dir, status="FAILED", code=error.failure_code,
        message=str(error), details=error.details)


def _record_mutex_guard_terminal(report_dir: Path, error,
                                 risk: RiskLevel) -> Path:
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    details = error.as_dict()
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.mutex_guard",
        selector_version=selector_version, risk=risk,
        environment={
            "mode": "mutex_guard_rejection",
            "lock_namespace": details.get("lock_namespace"),
            "mutex_denied_reason": details.get("mutex_denied_reason"),
        },
    )
    journal.emit(
        "mutex_guard", "FAILED", str(error),
        lock_namespace=details.get("lock_namespace"),
        mutex_denied_reason=details.get("mutex_denied_reason"),
        winerror=details.get("winerror"),
        business_data_modified=False, safe_to_retry=True)
    journal.emit(
        "run_finished", "FAILED", str(error),
        failure_code=error.failure_code,
        lock_namespace=details.get("lock_namespace"),
        mutex_denied_reason=details.get("mutex_denied_reason"),
        winerror=details.get("winerror"),
        business_data_modified=False, safe_to_retry=True)
    return journal.path


def _timed_wait(name: str):
    """Measure bounded waits without changing their calls or control flow."""
    def decorate(function):
        def wrapped(*args, **kwargs):
            started = time.monotonic()
            outcome = "FAILED"
            try:
                result = function(*args, **kwargs)
                outcome = "OK"
                return result
            finally:
                journal = kwargs.get("journal") or _ACTIVE_RUN_JOURNAL
                if journal is not None:
                    journal.emit(
                        "wait_timing", outcome, name,
                        wait=name,
                        wait_elapsed_ms=round(
                            (time.monotonic() - started) * 1000, 3),
                        timeout_seconds=kwargs.get("timeout"),
                    )
        return wrapped
    return decorate


def load_one_record(path: Path, sheet_name: str | None = None) -> RequisitionRecord:
    """Read exactly one business row without modifying the workbook."""
    from openpyxl import load_workbook

    if not path.is_file():
        raise FileNotFoundError(f"输入文件不存在: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"工作表不存在: {sheet_name}; 可选={workbook.sheetnames}")
            sheet = workbook[sheet_name]
        else:
            visible = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
            if len(visible) != 1:
                raise ValueError(
                    f"未指定 --sheet 且可见工作表数量为 {len(visible)}: "
                    f"{[ws.title for ws in visible]}")
            sheet = visible[0]

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("工作表为空")
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        duplicates = sorted({name for name in headers if name and headers.count(name) > 1})
        if duplicates:
            raise ValueError(f"表头重复: {duplicates}")
        missing = [name for name in REQUIRED_HEADERS if name not in headers]
        if missing:
            raise ValueError(f"缺少列: {', '.join(missing)}")

        business_rows = [row for row in rows[1:] if any(value not in (None, "") for value in row)]
        if len(business_rows) != 1:
            raise ValueError(
                f"首个写模板只允许恰好 1 条数据，实际 {len(business_rows)} 条")
        row = business_rows[0]
        mapping = {
            header: row[index] if index < len(row) else None
            for index, header in enumerate(headers) if header
        }
        return RequisitionRecord.from_mapping(mapping)
    finally:
        workbook.close()


def latest_recording(directory: Path) -> Path | None:
    candidates = list(directory.rglob("*.jsonl")) if directory.is_dir() else []
    return max(candidates, key=lambda item: item.stat().st_mtime) if candidates else None


def _display_value(value):
    if isinstance(value, dict):
        return str(value.get("display", ""))
    return str(value or "")


def align_recording(path: Path, record: RequisitionRecord) -> dict:
    """Compare a diagnostic recording with the input contract."""
    rows = []
    with path.open("r", encoding="utf-8") as stream:
        for line_number, line in enumerate(stream, 1):
            if line.strip():
                try:
                    rows.append(json.loads(line))
                except json.JSONDecodeError as exc:
                    raise ValueError(f"录制 JSONL 第 {line_number} 行损坏: {exc}") from exc

    entered = []
    for row in rows:
        if row.get("event_type") == "SET_VALUE":
            entered.append(_display_value(row.get("after")))
    wanted = {
        "单据类型": record.document_type,
        "申请人": record.applicant,
        "需求日期": record.required_date,
        "品号": record.item_no,
        "请购数量": record.quantity_text,
    }
    matched = {field: value in entered for field, value in wanted.items()}

    save_events = []
    warehouse_events = []
    document_numbers = []
    for row in rows:
        targets = [row.get("pre_target"), row.get("post_target"), row.get("target")]
        names = [str(target.get("name", "")) for target in targets if isinstance(target, dict)]
        if row.get("event_type") == "MOUSE_ACTION" and "保存" in names:
            save_events.append(row.get("seq"))
        if any("仓库" in name for name in names):
            warehouse_events.append(row.get("seq"))
        for name in names:
            if re.fullmatch(r"[A-Z]\d{6,}", name):
                document_numbers.append(name)
    return {
        "recording": str(path.resolve()),
        "session_id": rows[0].get("session_id") if rows else None,
        "events": len(rows),
        "entered_values": entered,
        "field_matches": matched,
        "all_input_fields_matched": all(matched.values()),
        "save_event_sequences": save_events,
        "document_numbers_seen_after_save": sorted(set(document_numbers)),
        "warehouse_event_sequences": warehouse_events,
        "warehouse_evidence": "RECORDED" if warehouse_events else "NOT_RECORDED",
    }


def semantic_plan(record: RequisitionRecord) -> list[dict]:
    return workflow_plan(
        "e10.requisition.create.end_to_end",
        warehouse_provided=record.warehouse is not None,
    )


def _print_plan(workflow_id: str, *, warehouse_provided: bool = False) -> None:
    plan = workflow_plan(workflow_id, warehouse_provided=warehouse_provided)
    for index, step in enumerate(plan, 1):
        print(
            f"[PLAN {index}/{len(plan)}] {step['step']}\n"
            f"  description={step['description']}\n"
            f"  risk={step['risk']}\n"
            f"  modifies_business_data={str(step['modifies_business_data']).lower()}\n"
            f"  retryable={str(step['retryable']).lower()}\n"
            f"  opens={step['opened_windows'] or '-'}\n"
            f"  closes={step['closed_windows'] or '-'}\n"
            f"  status={step['status']}"
            + (f" reason={step['reason']}" if step['reason'] else ""),
            flush=True,
        )


def _require_supported_warehouse(record: RequisitionRecord) -> None:
    if record.warehouse is not None:
        raise RuntimeError(
            "输入提供了仓库，但显式仓库选择尚未实现；拒绝静默使用 E10 默认仓库")


def _input_sha256(path: Path | None) -> str | None:
    if path is None:
        return None
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _stability(automation_id: str) -> str:
    value = str(automation_id or "").strip()
    if not value:
        return "missing"
    return "session_dynamic" if value.isdigit() else "candidate"


def _top_window(title: str, *, expected_pid: int | None = None) -> dict:
    from e10_recorder import enumerate_top_windows

    matches = [row for row in enumerate_top_windows(expected_pid)
               if (row.get("title", "").strip() == title
                   and _materialized_top_window(row))]
    if len(matches) != 1:
        raise RuntimeError(f"窗口 {title!r} 期望唯一，实际 {len(matches)} 个")
    return matches[0]


def _materialized_top_window(window: dict) -> bool:
    """Reject E10's visible-but-zero-rectangle transient/ghost top windows."""
    import win32api
    import win32con
    import win32gui

    try:
        hwnd = int(window["hwnd"])
        left, top, right, bottom = win32gui.GetWindowRect(hwnd)
        virtual_left = win32api.GetSystemMetrics(win32con.SM_XVIRTUALSCREEN)
        virtual_top = win32api.GetSystemMetrics(win32con.SM_YVIRTUALSCREEN)
        virtual_right = virtual_left + win32api.GetSystemMetrics(
            win32con.SM_CXVIRTUALSCREEN)
        virtual_bottom = virtual_top + win32api.GetSystemMetrics(
            win32con.SM_CYVIRTUALSCREEN)
        intersects_desktop = (
            right > virtual_left and left < virtual_right
            and bottom > virtual_top and top < virtual_bottom)
        return (bool(win32gui.IsWindowVisible(hwnd))
                and not win32gui.IsIconic(hwnd)
                and right > left and bottom > top
                and intersects_desktop)
    except Exception:
        return False


def _win32_window_candidates(title: str, expected_pid: int) -> list[dict]:
    """Enumerate all exact E10 candidates, including minimized windows."""
    import win32api
    import win32con
    import win32gui
    import win32process

    virtual_left = win32api.GetSystemMetrics(win32con.SM_XVIRTUALSCREEN)
    virtual_top = win32api.GetSystemMetrics(win32con.SM_YVIRTUALSCREEN)
    virtual_right = virtual_left + win32api.GetSystemMetrics(
        win32con.SM_CXVIRTUALSCREEN)
    virtual_bottom = virtual_top + win32api.GetSystemMetrics(
        win32con.SM_CYVIRTUALSCREEN)
    foreground = int(win32gui.GetForegroundWindow() or 0)
    rows = []

    def visit(hwnd, _extra):
        hwnd = int(hwnd)
        try:
            _thread_id, pid = win32process.GetWindowThreadProcessId(hwnd)
            text = str(win32gui.GetWindowText(hwnd) or "").strip()
            if int(pid) != int(expected_pid) or text != title:
                return True
            rect = tuple(int(value) for value in win32gui.GetWindowRect(hwnd))
            left, top, right, bottom = rect
            rows.append({
                "hwnd": hwnd,
                "pid": int(pid),
                "title": text,
                "visible": bool(win32gui.IsWindowVisible(hwnd)),
                "iconic": bool(win32gui.IsIconic(hwnd)),
                "foreground": hwnd == foreground,
                "rect": rect,
                "intersects_desktop": (
                    right > virtual_left and left < virtual_right
                    and bottom > virtual_top and top < virtual_bottom),
                "virtual_desktop": (
                    virtual_left, virtual_top, virtual_right, virtual_bottom),
            })
        except Exception:
            pass
        return True

    win32gui.EnumWindows(visit, None)
    return rows


def _recover_expected_window(title: str, expected_pid: int,
                             journal: RunJournal | None = None) -> dict | None:
    """Restore/raise one exact expected window without taskbar coordinates."""
    import win32con
    import win32gui

    before_candidates = _win32_window_candidates(title, expected_pid)
    plan = plan_window_recovery(
        before_candidates, title=title, expected_pid=expected_pid)
    if plan is None:
        return None

    applied = []
    if "RESTORE" in plan.actions:
        win32gui.ShowWindow(plan.hwnd, win32con.SW_RESTORE)
        applied.append("RESTORE")
        time.sleep(0.2)

    # Reinspect after restore: minimized windows can report sentinel/off-screen
    # rectangles until SW_RESTORE has completed.
    current = _win32_window_candidates(title, expected_pid)
    current_plan = plan_window_recovery(
        current, title=title, expected_pid=expected_pid)
    if current_plan is None:
        raise WindowRecoveryError(
            f"窗口 {title!r} 在恢复过程中消失: pid={expected_pid}")
    if "MOVE_TO_VISIBLE" in current_plan.actions:
        left, top, right, bottom = current_plan.before["rect"]
        virtual_left, virtual_top, _vr, _vb = current_plan.before[
            "virtual_desktop"]
        width = max(1, int(right) - int(left))
        height = max(1, int(bottom) - int(top))
        win32gui.SetWindowPos(
            current_plan.hwnd, win32con.HWND_TOP,
            int(virtual_left) + 32, int(virtual_top) + 32, width, height,
            win32con.SWP_SHOWWINDOW)
        applied.append("MOVE_TO_VISIBLE")

    # A temporary topmost toggle is the Win32 equivalent of selecting the
    # taskbar button, but remains PID/title scoped and is immediately released.
    flags = win32con.SWP_NOMOVE | win32con.SWP_NOSIZE | win32con.SWP_SHOWWINDOW
    win32gui.BringWindowToTop(current_plan.hwnd)
    win32gui.SetWindowPos(
        current_plan.hwnd, win32con.HWND_TOPMOST, 0, 0, 0, 0, flags)
    win32gui.SetWindowPos(
        current_plan.hwnd, win32con.HWND_NOTOPMOST, 0, 0, 0, 0, flags)
    try:
        win32gui.SetForegroundWindow(current_plan.hwnd)
    except Exception:
        pass
    applied.append("BRING_TO_FRONT")
    time.sleep(0.2)

    after_candidates = _win32_window_candidates(title, expected_pid)
    after_plan = plan_window_recovery(
        after_candidates, title=title, expected_pid=expected_pid)
    if after_plan is None or not _materialized_top_window({
            "hwnd": current_plan.hwnd}):
        raise WindowRecoveryError(
            f"窗口 {title!r} 执行恢复后仍未实体化: pid={expected_pid}")
    after = dict(after_plan.before)
    foreground_confirmed = bool(after.get("foreground"))
    if not foreground_confirmed:
        raise WindowRecoveryError(
            f"窗口 {title!r} 已恢复但未能置于前台: hwnd={current_plan.hwnd}")

    evidence = {
        "title": title,
        "pid": int(expected_pid),
        "hwnd": current_plan.hwnd,
        "actions": applied,
        "before": dict(plan.before),
        "after": after,
        "foreground_confirmed": foreground_confirmed,
        "permanent_topmost": False,
        "taskbar_coordinate_click": False,
    }
    if journal is not None:
        journal.emit(
            "window_recovery", "OK", "预期 E10 窗口已恢复并置于前台",
            **evidence, business_data_modified=False, safe_to_retry=True)
    return evidence


def _materialized(control) -> bool:
    try:
        rect = control.BoundingRectangle
        return (rect and rect.right > rect.left and rect.bottom > rect.top
                and not control.IsOffscreen)
    except Exception:
        return False


def _control_value(control) -> str | None:
    try:
        pattern = control.GetValuePattern()
        if pattern:
            value = pattern.Value
            if value is not None:
                return str(value).strip()
    except Exception:
        pass
    try:
        pattern = control.GetLegacyIAccessiblePattern()
        value = pattern.Value
        if value is not None:
            return str(value).strip()
    except Exception:
        pass
    try:
        value = control.Name
        return str(value).strip() if value is not None else None
    except Exception:
        return None


def _walk_controls(root, *, max_depth=16, limit=800):
    import uiautomation as auto

    controls = []
    for control, depth, _remaining in auto.WalkTree(
            root, getChildren=lambda item: item.GetChildren(),
            includeTop=False, maxDepth=max_depth):
        controls.append((control, depth))
        if len(controls) > limit:
            raise RuntimeError(f"控件遍历超过上限 {limit}")
    return controls


def _unique_descendant(root, control_class, key: str, **criteria):
    from e10_query import auto_unique

    control = auto_unique(root, key, control_class, timeout=3, **criteria)
    if control is None:
        raise RuntimeError(f"selector {key!r} 未命中")
    return control


def _immediate_materialized(scope, control_type: str):
    matches = []
    for child in scope.GetChildren():
        try:
            if child.ControlTypeName == control_type and _materialized(child):
                matches.append(child)
        except Exception:
            continue
    return matches


def _unique_native_edit(scope, label: str):
    native = []
    for child in scope.GetChildren():
        try:
            if (child.ControlTypeName == "EditControl"
                    and _materialized(child)
                    and int(child.NativeWindowHandle or 0) > 0):
                native.append(child)
        except Exception:
            continue
    if len(native) != 1:
        raise RuntimeError(f"{label}期望唯一原生 EDIT，实际 {len(native)}")
    return native[0]


def _field_scope(root, automation_id: str):
    import uiautomation as auto

    return _unique_descendant(
        root, auto.PaneControl, f"requisition.scope.{automation_id}",
        AutomationId=automation_id)


def _scope_values(scope) -> list[str]:
    values = []
    for control, _depth in _walk_controls(scope, max_depth=3, limit=30):
        if control.ControlTypeName == "EditControl" and _materialized(control):
            value = _control_value(control)
            if value:
                values.append(value)
    return values


@dataclass
class LiveForm:
    window: dict
    root: object
    document_type_scope: object
    applicant_scope: object
    required_date_scope: object
    grid_scope: object
    doc_no_edit: object


def attach_live_form(*, expected_pid: int | None = None) -> LiveForm:
    import uiautomation as auto

    window = _top_window("维护请购单", expected_pid=expected_pid)
    root = auto.ControlFromHandle(int(window["hwnd"]))
    if not root:
        raise RuntimeError("无法连接维护请购单 UIA 根节点")
    document_type_scope = _field_scope(root, "SelCtrDOC_ID")
    applicant_scope = _field_scope(root, "SelCtrSTAFF_ID")
    required_date_scope = _field_scope(root, "UDF_dtpUDF041")
    grid_scope = _field_scope(root, "grdREQUISITION_D")
    doc_no_edit = _unique_descendant(
        root, auto.EditControl, "requisition.doc_no", AutomationId="txtDOC_NO")
    return LiveForm(
        window, root, document_type_scope, applicant_scope,
        required_date_scope, grid_scope, doc_no_edit)


def live_preflight() -> dict:
    import uiautomation as auto

    form = attach_live_form()
    doc_no = _control_value(form.doc_no_edit) or ""
    if doc_no:
        raise RuntimeError(f"当前维护窗口不是空白新单，单号={doc_no!r}")

    button_counts = {}
    for key, scope in (("document_type", form.document_type_scope),
                       ("applicant", form.applicant_scope)):
        buttons = _immediate_materialized(scope, "ButtonControl")
        button_counts[key] = len(buttons)
        if len(buttons) != 2:
            raise RuntimeError(f"{key} 作用域期望恰好两个按钮，实际 {len(buttons)}")

    expected_cells = {}
    for column in ("品号", "请购数量"):
        name = f"{column} row -2147483647"
        cell = _unique_descendant(
            form.grid_scope, auto.DataItemControl,
            f"requisition.new_row.{column}", Name=name)
        expected_cells[column] = {
            "name": cell.Name,
            "value": _control_value(cell),
        }

    date_edits = [child for child in _immediate_materialized(
        form.required_date_scope, "EditControl")
        if str(child.AutomationId or "") == "txtDate"]
    if len(date_edits) != 1:
        raise RuntimeError(
            f"需求日期作用域内 txtDate 期望唯一，实际 {len(date_edits)}")

    return {
        "window": form.window,
        "doc_no": doc_no,
        "header_values": {
            "document_type": _scope_values(form.document_type_scope),
            "applicant": _scope_values(form.applicant_scope),
            "required_date": _scope_values(form.required_date_scope),
        },
        "lookup_button_counts": button_counts,
        "new_row_cells": expected_cells,
        "ready_for_fill": True,
    }


@_timed_wait("top_window")
def _wait_top_window(title: str, *, present=True, timeout=12.0,
                     expected_pid: int | None = None,
                     journal: RunJournal | None = None):
    from e10_recorder import enumerate_top_windows

    started = time.monotonic()
    deadline = time.monotonic() + timeout
    last_recovery_error = None
    recovery_attempted_for = set()
    while time.monotonic() < deadline:
        visible_title_matches = [
            row for row in enumerate_top_windows(expected_pid)
            if row.get("title", "").strip() == title]
        matches = [
            row for row in visible_title_matches
            if _materialized_top_window(row)]
        if present and len(matches) == 1:
            if expected_pid is not None:
                import win32gui

                hwnd = int(matches[0]["hwnd"])
                if int(win32gui.GetForegroundWindow() or 0) != hwnd:
                    _recover_expected_window(
                        title, expected_pid, journal=journal)
            return matches[0]
        if not present:
            remaining = (
                _win32_window_candidates(title, expected_pid)
                if expected_pid is not None else visible_title_matches)
            if not remaining:
                return None
            if len(remaining) > 1:
                raise RuntimeError(
                    f"等待关闭时窗口 {title!r} 仍命中 {len(remaining)} 个")
        if len(matches) > 1:
            raise RuntimeError(f"窗口 {title!r} 命中 {len(matches)} 个")

        # Give E10 a short normal creation grace period.  Then inspect all
        # exact-PID windows, including minimized/off-screen ones omitted by
        # the ordinary materialization filter.
        if (present and expected_pid is not None
                and time.monotonic() - started >= 0.6):
            candidates = _win32_window_candidates(title, expected_pid)
            signature = tuple(
                (row.get("hwnd"), row.get("visible"), row.get("iconic"),
                 tuple(row.get("rect") or ()))
                for row in candidates)
            if candidates and signature not in recovery_attempted_for:
                recovery_attempted_for.add(signature)
                try:
                    _recover_expected_window(
                        title, expected_pid, journal=journal)
                except WindowRecoveryError as exc:
                    last_recovery_error = exc
        time.sleep(0.2)
    state = "出现" if present else "关闭"
    if last_recovery_error is not None:
        raise TimeoutError(
            f"等待窗口 {title!r} {state} 超时；恢复失败: "
            f"{last_recovery_error}") from last_recovery_error
    raise TimeoutError(f"等待窗口 {title!r} {state} 超时")


def _query_button(scope, key: str):
    buttons = _immediate_materialized(scope, "ButtonControl")
    if len(buttons) != 2:
        raise RuntimeError(f"{key} 查询按钮作用域期望2个按钮，实际 {len(buttons)}")
    buttons.sort(key=lambda item: item.BoundingRectangle.left)
    if buttons[0].BoundingRectangle.left == buttons[1].BoundingRectangle.left:
        raise RuntimeError(f"{key} 两个按钮无法按相对顺序区分")
    return buttons[0]


@_timed_wait("exact_data_item")
def _find_exact_data_item(root, expected: str, *, timeout=8.0):
    deadline = time.monotonic() + timeout
    last_values = []
    while time.monotonic() < deadline:
        matches = []
        values = []
        for control, _depth in _walk_controls(root, max_depth=12, limit=800):
            if control.ControlTypeName != "DataItemControl" or not _materialized(control):
                continue
            value = _control_value(control) or ""
            values.append({"name": control.Name, "value": value})
            if value == expected:
                matches.append(control)
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise RuntimeError(
                f"查询结果值 {expected!r} 命中 {len(matches)} 个单元格，拒绝选择")
        last_values = values[:40]
        time.sleep(0.3)
    raise RuntimeError(
        f"查询结果中未找到精确值 {expected!r}；候选={last_values}")


def _select_lookup(scope, *, scope_key: str, dialog_title: str,
                   search_column: str, search_value: str,
                   journal: RunJournal | None = None):
    import uiautomation as auto
    from e10_query import auto_set_edit_value, auto_unique_in_pid

    button = _query_button(scope, scope_key)
    button.Click()
    expected_pid = int(scope.ProcessId)
    window = _wait_top_window(
        dialog_title, timeout=12, expected_pid=expected_pid,
        journal=journal)
    dialog = auto.ControlFromHandle(int(window["hwnd"]))
    pid = int(window["pid"])

    combo = _unique_descendant(
        dialog, auto.ComboBoxControl, f"{scope_key}.lookup.column",
        AutomationId="cmboxSimpleColumns")
    combo.Click()
    column_item = auto_unique_in_pid(
        pid, f"{scope_key}.lookup.column.{search_column}",
        auto.ListItemControl, timeout=5, Name=search_column)
    if column_item is None:
        raise RuntimeError(f"查询列选项 {search_column!r} 未命中")
    column_item.Click()

    keyword = _unique_descendant(
        dialog, auto.EditControl, f"{scope_key}.lookup.keyword",
        AutomationId="txtSearchKeyword")
    if not auto_set_edit_value(keyword, search_value):
        raise RuntimeError(f"查询值 {search_value!r} 写入后读回失败")
    search = _unique_descendant(
        dialog, auto.ButtonControl, f"{scope_key}.lookup.search",
        AutomationId="btnSimpleQuery")
    search.Click()
    result_cell = _find_exact_data_item(dialog, search_value, timeout=10)
    result_cell.Click()
    confirm = _unique_descendant(
        dialog, auto.ButtonControl, f"{scope_key}.lookup.confirm",
        AutomationId="btnOK")
    confirm.Click()
    _wait_top_window(
        dialog_title, present=False, timeout=10, expected_pid=expected_pid)
    return {
        "dialog": dialog_title,
        "search_column": search_column,
        "search_value": search_value,
        "result_cell": result_cell.Name,
    }


def _set_required_date(form: LiveForm, value: str):
    import uiautomation as auto
    from e10_query import auto_set_edit_value

    edit = _unique_descendant(
        form.required_date_scope, auto.EditControl,
        "requisition.required_date.edit", AutomationId="txtDate")
    if not auto_set_edit_value(edit, value):
        raise RuntimeError(f"需求日期 {value!r} 写入后读回失败")
    readback = _control_value(edit)
    if readback != value:
        raise RuntimeError(f"需求日期读回不一致: expected={value!r} actual={readback!r}")
    return readback


def _set_correlation_field(form: LiveForm, value: str) -> str:
    """Write the external request number into E10's queryable 备注 field."""
    from e10_query import auto_set_edit_value

    expected = correlation_value(value)
    scope = _field_scope(form.root, "selCtrREMARK")
    editor = _unique_native_edit(scope, "备注作用域")
    if not auto_set_edit_value(editor, expected):
        raise RuntimeError("业务请求号写入备注后读回失败")
    actual = _control_value(editor) or ""
    if actual != expected:
        raise RuntimeError(
            f"备注关联号读回不一致: expected={expected!r} actual={actual!r}")
    return actual


def _grid_cell(grid, column: str, row: int):
    import uiautomation as auto

    name = f"{column} row {row}"
    return _unique_descendant(
        grid, auto.DataItemControl, f"requisition.grid.{name}", Name=name)


def _top_windows_with_titles(titles: tuple[str, ...], *,
                             expected_pid: int | None = None) -> list[dict]:
    from e10_recorder import enumerate_top_windows

    wanted = set(titles)
    return [row for row in enumerate_top_windows(expected_pid)
            if (row.get("title", "").strip() in wanted
                and _materialized_top_window(row))]


@_timed_wait("item_lookup")
def _wait_item_lookup(*, timeout=10.0, expected_pid: int | None = None,
                      journal: RunJournal | None = None) -> dict:
    """Accept only the normal item lookup; asset lookup is a different action."""
    deadline = time.monotonic() + timeout
    watched = (ITEM_LOOKUP_TITLE, *UNEXPECTED_ITEM_LOOKUP_TITLES, "提示")
    started = time.monotonic()
    recovery_attempted_for = set()
    last_recovery_error = None
    while time.monotonic() < deadline:
        matches = _top_windows_with_titles(
            watched, expected_pid=expected_pid)
        normal = [row for row in matches if row["title"] == ITEM_LOOKUP_TITLE]
        if len(normal) == 1 and len(matches) == 1:
            if expected_pid is not None:
                import win32gui

                if int(win32gui.GetForegroundWindow() or 0) != int(
                        normal[0]["hwnd"]):
                    _recover_expected_window(
                        ITEM_LOOKUP_TITLE, expected_pid, journal=journal)
            return normal[0]
        if len(normal) > 1:
            raise RuntimeError(f"窗口 {ITEM_LOOKUP_TITLE!r} 命中多个")
        unexpected = [row["title"] for row in matches
                      if row["title"] != ITEM_LOOKUP_TITLE]
        if unexpected:
            raise RuntimeError(
                f"打开品号查询时出现非预期窗口 {unexpected}；"
                "拒绝把其他查询或确认框当成品号查询")
        if expected_pid is not None and time.monotonic() - started >= 0.6:
            candidates = _win32_window_candidates(
                ITEM_LOOKUP_TITLE, expected_pid)
            signature = tuple(
                (row.get("hwnd"), row.get("visible"), row.get("iconic"),
                 tuple(row.get("rect") or ()))
                for row in candidates)
            if candidates and signature not in recovery_attempted_for:
                recovery_attempted_for.add(signature)
                try:
                    _recover_expected_window(
                        ITEM_LOOKUP_TITLE, expected_pid, journal=journal)
                except WindowRecoveryError as exc:
                    # E10 creates a hidden titled shell before swapping to the
                    # visible item dialog.  Observe that transition during the
                    # bounded wait; only a shell that remains hidden is failure.
                    last_recovery_error = exc
        time.sleep(0.2)
    if last_recovery_error is not None:
        raise TimeoutError(
            f"等待窗口 {ITEM_LOOKUP_TITLE!r} 出现超时；"
            f"最后恢复证据: {last_recovery_error}") from last_recovery_error
    raise TimeoutError(f"等待窗口 {ITEM_LOOKUP_TITLE!r} 出现超时")


def dismiss_unexpected_item_lookup() -> dict:
    """Explicit fail-closed recovery; only clicks Cancel on a known wrong dialog."""
    import uiautomation as auto

    matches = _top_windows_with_titles(UNEXPECTED_ITEM_LOOKUP_TITLES)
    if len(matches) != 1:
        raise RuntimeError(
            "非预期品号查询窗口期望恰好一个，"
            f"实际 {[(row.get('title'), row.get('hwnd')) for row in matches]}")
    window = matches[0]
    root = auto.ControlFromHandle(int(window["hwnd"]))
    cancel = _unique_descendant(
        root, auto.ButtonControl, "unexpected_item_lookup.cancel",
        AutomationId="btnCancel")
    cancel.Click()
    _wait_top_window(
        window["title"], present=False, timeout=8,
        expected_pid=int(window["pid"]))
    return {"dismissed": window["title"], "action": "btnCancel"}


def _activate_item_editor(form: LiveForm):
    import uiautomation as auto

    try:
        return _unique_descendant(
            form.grid_scope, auto.PaneControl,
            "requisition.item.editor_scope", AutomationId="SelCtrITEM_ID_CODE")
    except RuntimeError:
        _grid_cell(form.grid_scope, "品号", -2147483647).Click()

    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        try:
            return _unique_descendant(
                form.grid_scope, auto.PaneControl,
                "requisition.item.editor_scope", AutomationId="SelCtrITEM_ID_CODE")
        except RuntimeError:
            time.sleep(0.2)
    raise RuntimeError("点击品号新项目行后未出现 SelCtrITEM_ID_CODE 编辑器")


def _open_item_lookup(form: LiveForm,
                      journal: RunJournal | None = None) -> dict:
    """Open the ordinary item query through the editor's documented F2 action."""
    from e10_query import auto_set_edit_value

    editor_scope = _activate_item_editor(form)
    editor = _unique_native_edit(editor_scope, "品号编辑器")
    prior_value = _control_value(editor) or ""
    if prior_value:
        if not auto_set_edit_value(editor, ""):
            raise RuntimeError("品号编辑器残留值无法清空，拒绝继续")
        if _control_value(editor):
            raise RuntimeError("品号编辑器清空后读回仍非空，拒绝继续")

    editor.SetFocus()
    editor.SendKeys(ITEM_LOOKUP_HOTKEY)
    window = _wait_item_lookup(
        timeout=10, expected_pid=int(form.window["pid"]), journal=journal)
    return {
        "method": "active_item_editor_hotkey",
        "hotkey": "F2",
        "cleared_residual_value": bool(prior_value),
        "dialog": window["title"],
    }


def _row_index_from_cell_name(name: str) -> int:
    match = re.search(r"\brow\s+(-?\d+)\s*$", str(name or ""))
    if not match:
        raise RuntimeError(f"查询结果单元格缺少稳定行号: {name!r}")
    return int(match.group(1))


@_timed_wait("item_dialog_ready")
def _wait_item_dialog_ready(*, timeout=12.0, expected_pid: int | None = None,
                            journal: RunJournal | None = None):
    """Wait through E10's shell-handle swap until the search controls exist."""
    import uiautomation as auto

    deadline = time.monotonic() + timeout
    observed = []
    recovery_attempted_for = set()
    last_recovery_error = None
    while time.monotonic() < deadline:
        windows = _top_windows_with_titles(
            (ITEM_LOOKUP_TITLE,), expected_pid=expected_pid)
        if len(windows) > 1:
            raise RuntimeError(
                f"实体化的 {ITEM_LOOKUP_TITLE!r} 窗口命中多个: "
                f"{[(row['hwnd'], row['pid']) for row in windows]}")
        if windows:
            window = windows[0]
            observed.append(int(window["hwnd"]))
            dialog = auto.ControlFromHandle(int(window["hwnd"]))
            combo = auto.ComboBoxControl(
                searchFromControl=dialog, searchDepth=30, foundIndex=1,
                AutomationId="cmboxSimpleColumns")
            if combo.Exists(0.5, 0.1):
                duplicate = auto.ComboBoxControl(
                    searchFromControl=dialog, searchDepth=30, foundIndex=2,
                    AutomationId="cmboxSimpleColumns")
                if duplicate.Exists(0.2, 0.1):
                    raise RuntimeError("item.lookup.column 命中多个控件")
                if _materialized(combo):
                    return window, dialog, combo
        if expected_pid is not None and not windows:
            candidates = _win32_window_candidates(
                ITEM_LOOKUP_TITLE, expected_pid)
            signature = tuple(
                (row.get("hwnd"), row.get("visible"), row.get("iconic"),
                 tuple(row.get("rect") or ()))
                for row in candidates)
            if candidates and signature not in recovery_attempted_for:
                recovery_attempted_for.add(signature)
                try:
                    _recover_expected_window(
                        ITEM_LOOKUP_TITLE, expected_pid, journal=journal)
                except WindowRecoveryError as exc:
                    last_recovery_error = exc
        time.sleep(0.2)
    recovery_note = (
        f"；最后恢复证据={last_recovery_error}"
        if last_recovery_error is not None else "")
    raise TimeoutError(
        "查询品号窗口已出现但搜索控件未就绪；"
        f"观察到的实体窗口句柄={observed[-10:]}{recovery_note}")


def _select_open_item_dialog(form: LiveForm, item_no: str,
                             journal: RunJournal | None = None):
    """Complete item lookup after ``_select_item`` has opened the dialog."""
    import uiautomation as auto
    from e10_query import auto_set_edit_value, auto_unique_in_pid

    expected_pid = int(form.window["pid"])
    window, dialog, combo = _wait_item_dialog_ready(
        timeout=12, expected_pid=expected_pid, journal=journal)
    pid = int(window["pid"])
    combo.Click()
    column_item = auto_unique_in_pid(
        pid, "item.lookup.column.品号", auto.ListItemControl,
        timeout=5, Name="品号")
    if column_item is None:
        raise RuntimeError("品号查询列未命中")
    column_item.Click()
    keyword = _unique_descendant(
        dialog, auto.EditControl, "item.lookup.keyword",
        AutomationId="txtSearchKeyword")
    if not auto_set_edit_value(keyword, item_no):
        raise RuntimeError(f"品号 {item_no!r} 写入后读回失败")
    search = _unique_descendant(
        dialog, auto.ButtonControl, "item.lookup.search",
        AutomationId="btnSimpleQuery")
    search.Click()
    result = _find_exact_data_item(dialog, item_no, timeout=10)
    row_index = _row_index_from_cell_name(result.Name)
    selection = _unique_descendant(
        dialog, auto.DataItemControl, "item.lookup.selection",
        Name=f"选定 row {row_index}")
    selection.Click()
    confirm = _unique_descendant(
        dialog, auto.ButtonControl, "item.lookup.confirm",
        AutomationId="btnOK")
    confirm.Click()
    _wait_top_window(
        ITEM_LOOKUP_TITLE, present=False, timeout=10,
        expected_pid=expected_pid)
    selected = _grid_cell(form.grid_scope, "品号", 0)
    readback = _control_value(selected)
    if readback != item_no:
        raise RuntimeError(f"品号回填不一致: expected={item_no!r} actual={readback!r}")
    return {
        "result_cell": result.Name,
        "selection_cell": selection.Name,
        "readback": readback,
    }


def _set_item_from_active_editor(form: LiveForm, item_no: str,
                                 journal: RunJournal | None = None):
    opening = _open_item_lookup(form, journal=journal)
    selected = _select_open_item_dialog(form, item_no, journal=journal)
    return {"opening": opening, **selected}


def _set_quantity(form: LiveForm, value: str):
    import uiautomation as auto
    from e10_query import auto_set_edit_value

    cell = _grid_cell(form.grid_scope, "请购数量", 0)
    cell.Click()
    deadline = time.monotonic() + 5
    editor = None
    while time.monotonic() < deadline:
        try:
            scope = _unique_descendant(
                form.grid_scope, auto.PaneControl,
                "requisition.quantity.editor_scope",
                AutomationId="colREQUISITION_QTYSelectWindowEditor")
        except RuntimeError:
            time.sleep(0.2)
            continue
        # The same WinForms EDIT is also exposed as a duplicate proxy with no
        # native handle. Select by native child status inside the stable scope,
        # never by the session-dynamic numeric AutomationId.
        try:
            editor = _unique_native_edit(scope, "数量编辑器作用域")
            break
        except RuntimeError as exc:
            if "实际 0" not in str(exc):
                raise
        time.sleep(0.2)
    if editor is None:
        raise RuntimeError("请购数量单元格未出现唯一数据编辑器")
    if not auto_set_edit_value(editor, value):
        raise RuntimeError(f"请购数量 {value!r} 写入后读回失败")
    editor.SendKeys("{ENTER}")
    time.sleep(0.5)
    cell = _grid_cell(form.grid_scope, "请购数量", 0)
    readback = (_control_value(cell) or "").replace(",", "")
    if readback != value:
        raise RuntimeError(f"请购数量回填不一致: expected={value!r} actual={readback!r}")
    return readback


def _capture_screen(directory: Path, tag: str) -> Path:
    from PIL import ImageGrab

    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{datetime.now():%H%M%S}_{tag}.png"
    ImageGrab.grab(all_screens=True).save(path)
    return path


def fill_live_form(record: RequisitionRecord, report_dir: Path) -> dict:
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.create.fill",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={"account_set": "FRKTEST", "mode": "fill_without_save"},
    )
    try:
        _require_supported_warehouse(record)
        preflight = live_preflight()
        journal.emit("preflight", "OK", "空白新单前置检查通过",
                     step="preflight", **preflight)
        form = attach_live_form()

        doc_evidence = _select_lookup(
            form.document_type_scope,
            scope_key="document_type",
            dialog_title="查询单据类型",
            search_column="单据类型",
            search_value=record.document_type,
            journal=journal,
        )
        doc_values = _scope_values(form.document_type_scope)
        if (record.document_type not in doc_values
                and record.document_name not in doc_values):
            raise RuntimeError(f"单据类型回填未命中预期，读回={doc_values}")
        journal.emit("field", "OK", "单据类型已选择",
                     step="select_document_type",
                     field="单据类型", values=doc_values, evidence=doc_evidence)

        applicant_evidence = _select_lookup(
            form.applicant_scope,
            scope_key="applicant",
            dialog_title="查询员工",
            search_column="员工姓名",
            search_value=record.applicant,
            journal=journal,
        )
        applicant_values = _scope_values(form.applicant_scope)
        if record.applicant not in applicant_values:
            raise RuntimeError(f"申请人回填未命中预期，读回={applicant_values}")
        journal.emit("field", "OK", "申请人已选择",
                     step="select_applicant", field="申请人", values=applicant_values,
                     evidence=applicant_evidence)

        required_date = _set_required_date(form, record.required_date)
        journal.emit("field", "OK", "需求日期已写入",
                     step="set_required_date", field="需求日期", value=required_date)

        item_evidence = _set_item_from_active_editor(
            form, record.item_no, journal=journal)
        journal.emit("field", "OK", "品号已选择",
                     step="select_item", field="品号", value=record.item_no,
                     evidence=item_evidence)

        quantity = _set_quantity(form, record.quantity_text)
        journal.emit("field", "OK", "请购数量已写入",
                     step="set_quantity", field="请购数量", value=quantity)
        journal.emit("step", "SKIPPED", "输入未提供仓库；记录 E10 自动默认值",
                     step="verify_warehouse",
                     reason="not_in_input_e10_auto_default")

        shot = _capture_screen(report_dir / f"shots_{journal.run_id[:8]}", "filled")
        journal.emit("screenshot", "OK", "填写完成截图", path=str(shot))
        result = {
            "status": "FILLED_NOT_SAVED",
            "journal": str(journal.path),
            "screenshot": str(shot),
            "record": record.as_business_mapping(),
            "business_data_modified": False,
            "safe_to_retry": False,
        }
        journal.emit("run_finished", "FILLED_NOT_SAVED", "字段填写并读回完成，尚未保存",
                     business_data_modified=False, safe_to_retry=False,
                     record=record.as_business_mapping())
        return result
    except Exception as exc:
        journal.emit("run_finished", "FAILED", str(exc),
                     business_data_modified=False, safe_to_retry=False,
                     error_type=type(exc).__name__)
        raise


def resume_live_form(record: RequisitionRecord, report_dir: Path) -> dict:
    """Resume the exact draft left by a failed fill before any Save click."""
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.create.resume_fill",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={"account_set": "FRKTEST", "mode": "resume_fill_without_save"},
    )
    try:
        _require_supported_warehouse(record)
        form = attach_live_form()
        state = _validate_resume_state(form, record)
        doc_no = state["doc_no"]
        journal.emit(
            "resume_precondition", "OK", "已确认是本轮表头已填、表体未填的预分配新单",
            step="resume_precondition", **state)

        item_evidence = _set_item_from_active_editor(
            form, record.item_no, journal=journal)
        journal.emit("field", "OK", "品号已写入并读回",
                     step="select_item", field="品号", value=record.item_no,
                     evidence=item_evidence)
        quantity = _set_quantity(form, record.quantity_text)
        journal.emit("field", "OK", "请购数量已写入并读回",
                     step="set_quantity", field="请购数量", value=quantity)
        journal.emit("step", "SKIPPED", "输入未提供仓库；记录 E10 自动默认值",
                     step="verify_warehouse",
                     reason="not_in_input_e10_auto_default")
        shot = _capture_screen(report_dir / f"shots_{journal.run_id[:8]}", "filled")
        journal.emit("screenshot", "OK", "填写完成截图", path=str(shot))
        result = {
            "status": "FILLED_NOT_SAVED",
            "doc_no": doc_no,
            "journal": str(journal.path),
            "screenshot": str(shot),
            "record": record.as_business_mapping(),
            "business_data_modified": False,
            "safe_to_retry": False,
        }
        journal.emit("run_finished", "FILLED_NOT_SAVED", "续跑填写完成，尚未保存",
                     business_data_modified=False, safe_to_retry=False,
                     doc_no=doc_no, record=record.as_business_mapping())
        return result
    except Exception as exc:
        journal.emit("run_finished", "FAILED", str(exc),
                     business_data_modified=False, safe_to_retry=False,
                     error_type=type(exc).__name__)
        raise


def _validate_resume_state(form: LiveForm, record: RequisitionRecord) -> dict:
    doc_no = _control_value(form.doc_no_edit) or ""
    document_values = _scope_values(form.document_type_scope)
    applicant_values = _scope_values(form.applicant_scope)
    date_values = _scope_values(form.required_date_scope)
    if not doc_no:
        raise RuntimeError("续跑要求 E10 已为当前新单预分配单号")
    if (record.document_type not in document_values
            and record.document_name not in document_values):
        raise RuntimeError(f"续跑单据类型不匹配: {document_values}")
    if record.applicant not in applicant_values:
        raise RuntimeError(f"续跑申请人不匹配: {applicant_values}")
    if record.required_date not in date_values:
        raise RuntimeError(f"续跑需求日期不匹配: {date_values}")
    _grid_cell(form.grid_scope, "品号", -2147483647)
    from e10_query import auto_collect
    import uiautomation as auto
    material_rows = auto_collect(
        form.grid_scope, auto.DataItemControl, limit=3, timeout=0.3,
        Name="品号 row 0")
    if material_rows:
        raise RuntimeError("续跑表体已存在 row 0，拒绝再次执行品号选择")
    return {
        "doc_no": doc_no,
        "document_values": document_values,
        "applicant_values": applicant_values,
        "date_values": date_values,
    }


def resume_open_item_dialog(record: RequisitionRecord, report_dir: Path) -> dict:
    """Continue an interrupted run from one already-open normal item dialog."""
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.create.resume_open_item",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={"account_set": "FRKTEST", "mode": "resume_open_item_without_save"},
    )
    try:
        _require_supported_warehouse(record)
        form = attach_live_form()
        state = _validate_resume_state(form, record)
        window = _wait_top_window(ITEM_LOOKUP_TITLE, timeout=3)
        journal.emit(
            "resume_precondition", "OK",
            "表头与空白表体匹配，且普通品号查询窗口已实体化",
            step="resume_precondition", **state,
            item_lookup_hwnd=window["hwnd"])
        item_evidence = _select_open_item_dialog(form, record.item_no)
        journal.emit("field", "OK", "品号已选择并读回",
                     step="select_item", field="品号", value=record.item_no,
                     evidence=item_evidence)
        quantity = _set_quantity(form, record.quantity_text)
        journal.emit("field", "OK", "请购数量已写入并读回",
                     step="set_quantity", field="请购数量", value=quantity)
        journal.emit("step", "SKIPPED", "输入未提供仓库；记录 E10 自动默认值",
                     step="verify_warehouse",
                     reason="not_in_input_e10_auto_default")
        shot = _capture_screen(report_dir / f"shots_{journal.run_id[:8]}", "filled")
        journal.emit("screenshot", "OK", "填写完成截图", path=str(shot))
        result = {
            "status": "FILLED_NOT_SAVED",
            "doc_no": state["doc_no"],
            "journal": str(journal.path),
            "screenshot": str(shot),
            "record": record.as_business_mapping(),
            "business_data_modified": False,
            "safe_to_retry": False,
        }
        journal.emit(
            "run_finished", "FILLED_NOT_SAVED", "查询窗续跑填写完成，尚未保存",
            business_data_modified=False, safe_to_retry=False,
            doc_no=state["doc_no"], record=record.as_business_mapping())
        return result
    except Exception as exc:
        journal.emit("run_finished", "FAILED", str(exc),
                     business_data_modified=False, safe_to_retry=False,
                     error_type=type(exc).__name__)
        raise


def resume_quantity(record: RequisitionRecord, report_dir: Path) -> dict:
    """Continue only when row 0 already contains the exact requested item."""
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.create.resume_quantity",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={"account_set": "FRKTEST", "mode": "resume_quantity_without_save"},
    )
    try:
        _require_supported_warehouse(record)
        form = attach_live_form()
        doc_no = _control_value(form.doc_no_edit) or ""
        if not doc_no:
            raise RuntimeError("数量续跑要求预分配单号非空")
        item_cell = _grid_cell(form.grid_scope, "品号", 0)
        item_readback = _control_value(item_cell)
        if item_readback != record.item_no:
            raise RuntimeError(
                f"数量续跑品号不匹配: expected={record.item_no!r} "
                f"actual={item_readback!r}")
        quantity_cell = _grid_cell(form.grid_scope, "请购数量", 0)
        before_quantity = (_control_value(quantity_cell) or "").replace(",", "")
        if before_quantity == record.quantity_text:
            quantity = before_quantity
        elif before_quantity in ("", "0", "0.0", "0.00"):
            quantity = _set_quantity(form, record.quantity_text)
        else:
            raise RuntimeError(
                f"数量续跑发现非零已有值 {before_quantity!r}，拒绝覆盖")
        journal.emit(
            "resume_precondition", "OK", "已确认 row 0 品号精确匹配",
            step="resume_precondition", doc_no=doc_no, item=item_readback,
            before_quantity=before_quantity)
        journal.emit("field", "OK", "请购数量已写入并读回",
                     step="set_quantity", field="请购数量", value=quantity)
        journal.emit("step", "SKIPPED", "输入未提供仓库；记录 E10 自动默认值",
                     step="verify_warehouse",
                     reason="not_in_input_e10_auto_default")
        shot = _capture_screen(report_dir / f"shots_{journal.run_id[:8]}", "filled")
        journal.emit("screenshot", "OK", "填写完成截图", path=str(shot))
        result = {
            "status": "FILLED_NOT_SAVED",
            "doc_no": doc_no,
            "journal": str(journal.path),
            "screenshot": str(shot),
            "record": record.as_business_mapping(),
            "business_data_modified": False,
            "safe_to_retry": False,
        }
        journal.emit(
            "run_finished", "FILLED_NOT_SAVED", "数量续跑填写完成，尚未保存",
            business_data_modified=False, safe_to_retry=False,
            doc_no=doc_no, record=record.as_business_mapping())
        return result
    except Exception as exc:
        journal.emit("run_finished", "FAILED", str(exc),
                     business_data_modified=False, safe_to_retry=False,
                     error_type=type(exc).__name__)
        raise


def validate_filled_form(record: RequisitionRecord, report_dir: Path, *,
                         correlation: str | None = None) -> dict:
    """Run E10 validation once after exact row readback; never clicks Save."""
    import uiautomation as auto
    from e10_recorder import enumerate_top_windows

    form = attach_live_form()
    doc_no = _control_value(form.doc_no_edit) or ""
    item = _control_value(_grid_cell(form.grid_scope, "品号", 0))
    quantity = (_control_value(
        _grid_cell(form.grid_scope, "请购数量", 0)) or "").replace(",", "")
    if not doc_no or item != record.item_no or quantity != record.quantity_text:
        raise RuntimeError(
            f"校验前读回不一致: doc_no={doc_no!r}, item={item!r}, "
            f"quantity={quantity!r}")
    correlation_before = None
    if correlation is not None:
        correlation_before = _control_value(_unique_native_edit(
            _field_scope(form.root, "selCtrREMARK"), "备注作用域")) or ""
        if correlation_before != correlation_value(correlation):
            raise RuntimeError(
                "校验前备注关联号读回不一致: "
                f"expected={correlation!r} actual={correlation_before!r}")
    before = {int(row["hwnd"]) for row in enumerate_top_windows()}
    button = _unique_descendant(
        form.root, auto.ButtonControl, "requisition.validate", Name="校验")
    button.Click()
    time.sleep(2)
    unexpected = []
    for row in enumerate_top_windows():
        if (int(row["hwnd"]) in before
                or not _materialized_top_window(row)
                or not str(row.get("title", "")).strip()):
            continue
        unexpected.append({key: row.get(key) for key in ("hwnd", "pid", "title")})
    shot = _capture_screen(report_dir / "validation", "validated")
    if unexpected:
        raise RuntimeError(
            f"E10 校验后出现新窗口，拒绝保存: {unexpected}; screenshot={shot}")
    item_after = _control_value(_grid_cell(form.grid_scope, "品号", 0))
    quantity_after = (_control_value(
        _grid_cell(form.grid_scope, "请购数量", 0)) or "").replace(",", "")
    if item_after != record.item_no or quantity_after != record.quantity_text:
        raise RuntimeError("E10 校验后关键字段读回发生变化，拒绝保存")
    correlation_after = correlation_before
    if correlation is not None:
        correlation_after = _control_value(_unique_native_edit(
            _field_scope(form.root, "selCtrREMARK"), "备注作用域")) or ""
        if correlation_after != correlation_value(correlation):
            raise RuntimeError("E10 校验后备注关联号发生变化，拒绝保存")
    return {
        "status": "VALIDATED_NOT_SAVED",
        "doc_no": doc_no,
        "item": item_after,
        "quantity": quantity_after,
        "correlation_field": "备注" if correlation is not None else None,
        "correlation_value": correlation_after,
        "screenshot": str(shot),
        "unexpected_windows": [],
    }


def _post_close(hwnd: int) -> None:
    import win32con
    import win32gui

    win32gui.PostMessage(int(hwnd), win32con.WM_CLOSE, 0, 0)


@_timed_wait("hwnd_closed")
def _wait_hwnd_closed(hwnd: int, *, timeout=8.0) -> bool:
    import win32gui

    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if not win32gui.IsWindow(int(hwnd)):
            return True
        time.sleep(0.2)
    return not win32gui.IsWindow(int(hwnd))


def _classify_discardable_draft_row(record: RequisitionRecord,
                                    item: str | None,
                                    quantity: str | None) -> str:
    item_value = str(item or "").strip()
    quantity_value = str(quantity or "").replace(",", "").strip()
    if (item_value == record.item_no
            and quantity_value == record.quantity_text):
        return "filled"
    if not item_value and quantity_value in ("", "0", "0.0", "0.00"):
        return "header_only_with_blank_row"
    raise RuntimeError(
        "当前草稿 row 0 与本轮 XLSX 不一致，拒绝代替用户丢弃: "
        f"item={item_value!r}, quantity={quantity_value!r}")


def _discard_current_test_draft_impl(record: RequisitionRecord) -> dict:
    """Discard only the exact FRKTEST draft created during this test session."""
    import uiautomation as auto
    from e10_recorder import enumerate_top_windows

    form = attach_live_form()
    doc_no = _control_value(form.doc_no_edit) or ""
    try:
        item = _control_value(_grid_cell(form.grid_scope, "品号", 0))
        quantity = (_control_value(
            _grid_cell(form.grid_scope, "请购数量", 0)) or "").replace(",", "")
        draft_stage = _classify_discardable_draft_row(record, item, quantity)
        if draft_stage == "header_only_with_blank_row":
            document_values = _scope_values(form.document_type_scope)
            applicant_values = _scope_values(form.applicant_scope)
            date_values = _scope_values(form.required_date_scope)
            if ((record.document_type not in document_values
                 and record.document_name not in document_values)
                    or record.applicant not in applicant_values
                    or record.required_date not in date_values):
                raise RuntimeError(
                    "当前空白占位行草稿的表头与本轮 XLSX 不一致，拒绝丢弃")
    except RuntimeError as exc:
        if "未命中" not in str(exc):
            raise
        document_values = _scope_values(form.document_type_scope)
        applicant_values = _scope_values(form.applicant_scope)
        date_values = _scope_values(form.required_date_scope)
        if ((record.document_type not in document_values
             and record.document_name not in document_values)
                or record.applicant not in applicant_values
                or record.required_date not in date_values):
            raise RuntimeError(
                "当前无明细草稿的表头与本轮 XLSX 不一致，拒绝丢弃")
        item = None
        quantity = None
        draft_stage = "header_only"
    pid = int(form.window["pid"])
    edit_hwnd = int(form.window["hwnd"])
    baseline = {int(row["hwnd"]) for row in enumerate_top_windows()}
    _post_close(edit_hwnd)

    clicked_no = False
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline and not _wait_hwnd_closed(edit_hwnd, timeout=0.2):
        prompts = []
        for row in enumerate_top_windows():
            if (int(row.get("pid", 0)) != pid
                    or int(row["hwnd"]) in baseline
                    or not _materialized_top_window(row)):
                continue
            root = auto.ControlFromHandle(int(row["hwnd"]))
            no = auto.ButtonControl(
                searchFromControl=root, searchDepth=12,
                foundIndex=1, AutomationId="btnNo")
            if no.Exists(0.2, 0.1) and _materialized(no):
                prompts.append((row, no))
        if len(prompts) > 1:
            raise RuntimeError("关闭草稿时出现多个带‘否’按钮的确认窗")
        if len(prompts) == 1:
            prompts[0][1].Click()
            clicked_no = True
        time.sleep(0.2)
    if not _wait_hwnd_closed(edit_hwnd, timeout=1):
        raise RuntimeError("当前测试草稿未能安全关闭")

    browse = _top_window("浏览 - 维护请购单")
    browse_hwnd = int(browse["hwnd"])
    _post_close(browse_hwnd)
    if not _wait_hwnd_closed(browse_hwnd, timeout=8):
        raise RuntimeError("维护请购单浏览窗口未能关闭")
    return {
        "discarded_doc_no": doc_no,
        "discarded_stage": draft_stage,
        "discard_confirmation_no_clicked": clicked_no,
        "main_page_restored": True,
    }


def discard_current_test_draft(record: RequisitionRecord,
                               report_dir: Path) -> dict:
    """Evidence wrapper around the unchanged draft-discard actuator path."""
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.create.discard_draft",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={"account_set": "FRKTEST", "mode": "discard_unsaved_draft"},
    )
    started = time.monotonic()
    try:
        result = _discard_current_test_draft_impl(record)
        journal.emit(
            "run_finished", "DISCARDED", "精确匹配的未保存测试草稿已安全丢弃",
            business_data_modified=False, safe_to_retry=True,
            step_elapsed_ms=round((time.monotonic() - started) * 1000, 3),
            **result)
        return {**result, "journal": str(journal.path)}
    except Exception as exc:
        journal.emit(
            "run_finished", "FAILED", str(exc),
            business_data_modified=False, safe_to_retry=True,
            step_elapsed_ms=round((time.monotonic() - started) * 1000, 3),
            error_type=type(exc).__name__)
        raise


def _resolve_presave_orphans(record: RequisitionRecord, report_dir: Path, *,
                             request_no: str, dataset_epoch: str,
                             orphans: list[dict]) -> dict:
    """Resolve interrupted runs known to have stopped before write issuance."""
    from e10_recorder import enumerate_top_windows

    for orphan in orphans:
        fingerprint = str(orphan.get("payload_fingerprint") or "")
        if fingerprint:
            restored_fingerprint = fingerprint
        else:
            source = Path(str(orphan.get("input_file") or ""))
            expected_hash = str(orphan.get("input_hash") or "")
            if not source.is_file() or not expected_hash:
                raise RuntimeError(
                    "无终态运行缺少可校验的 input_file/input_hash，拒绝猜测草稿")
            actual_hash = _input_sha256(source)
            if actual_hash != expected_hash:
                raise RuntimeError(
                    "无终态运行的输入文件哈希已变化，拒绝用新内容处置旧草稿")
            restored_fingerprint = load_one_record(
                source, None).payload_fingerprint
        if restored_fingerprint != record.payload_fingerprint:
            raise RuntimeError(
                "无终态运行的载荷与本次输入不同，拒绝自动处理草稿")
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))[
            "_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.orphan_guard",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={
            "account_set": "FRKTEST", "mode": "orphan_presave_guard",
            "request_no": request_no, "dataset_epoch": dataset_epoch,
            "payload_fingerprint": record.payload_fingerprint,
        })
    edit_windows = [
        row for row in enumerate_top_windows()
        if str(row.get("title") or "").strip() == "维护请购单"
        and _materialized_top_window(row)
    ]
    if len(edit_windows) > 1:
        raise RuntimeError(
            f"无终态运行检查发现 {len(edit_windows)} 个请购编辑窗，拒绝猜测")
    cleanup = None
    if edit_windows:
        cleanup = _discard_current_test_draft_impl(record)
        resolution = "MATCHED_DRAFT_DISCARDED"
    else:
        resolution = "NO_OPEN_DRAFT"
    orphan_ids = [str(row.get("run_id")) for row in orphans]
    journal.emit(
        "startup_reconciliation", "OK",
        "无终态且未发出保存的运行已完成草稿检查",
        resolved_orphan_run_ids=orphan_ids,
        resolution=resolution, cleanup=cleanup,
        business_data_modified=False, safe_to_retry=True)
    journal.emit(
        "run_finished", "NOT_APPLIED", "保存前中断样本已完成草稿处置",
        resolved_orphan_run_ids=orphan_ids,
        resolution=resolution, cleanup=cleanup,
        business_data_modified=False, safe_to_retry=True)
    return {"journal": str(journal.path), "resolution": resolution,
            "resolved_orphan_run_ids": orphan_ids, "cleanup": cleanup}


def _fill_current_form(form: LiveForm, record: RequisitionRecord,
                       journal: RunJournal,
                       progress: ProgressHeartbeat | None = None, *,
                       correlation: str) -> dict:
    if progress:
        progress.begin_step("select_document_type")
    doc_evidence = _select_lookup(
        form.document_type_scope,
        scope_key="document_type",
        dialog_title="查询单据类型",
        search_column="单据类型",
        search_value=record.document_type,
        journal=journal,
    )
    doc_values = _scope_values(form.document_type_scope)
    if (record.document_type not in doc_values
            and record.document_name not in doc_values):
        raise RuntimeError(f"单据类型回填未命中预期，读回={doc_values}")
    journal.emit("field", "OK", "单据类型已选择",
                 step="select_document_type",
                 field="单据类型", values=doc_values, evidence=doc_evidence)
    if progress:
        progress.complete_step("select_document_type")

    if progress:
        progress.begin_step("select_applicant")
    applicant_evidence = _select_lookup(
        form.applicant_scope,
        scope_key="applicant",
        dialog_title="查询员工",
        search_column="员工姓名",
        search_value=record.applicant,
        journal=journal,
    )
    applicant_values = _scope_values(form.applicant_scope)
    if record.applicant not in applicant_values:
        raise RuntimeError(f"申请人回填未命中预期，读回={applicant_values}")
    journal.emit("field", "OK", "申请人已选择",
                 step="select_applicant", field="申请人", values=applicant_values,
                 evidence=applicant_evidence)
    if progress:
        progress.complete_step("select_applicant")

    if progress:
        progress.begin_step("set_required_date")
    required_date = _set_required_date(form, record.required_date)
    journal.emit("field", "OK", "需求日期已写入",
                 step="set_required_date", field="需求日期", value=required_date)
    if progress:
        progress.complete_step("set_required_date")
        progress.begin_step("set_correlation_field")
    correlation_readback = _set_correlation_field(form, correlation)
    journal.emit(
        "field", "OK", "业务请求号已写入备注并读回",
        step="set_correlation_field", field="备注",
        correlation_field="selCtrREMARK",
        correlation_value=correlation_readback,
        identity_strength=IdentityStrength.REQUEST_FIELD_MATCH.value)
    if progress:
        progress.complete_step("set_correlation_field")
        progress.begin_step("select_item")
    item_evidence = _set_item_from_active_editor(
        form, record.item_no, journal=journal)
    journal.emit("field", "OK", "品号已选择并读回",
                 step="select_item", field="品号", value=record.item_no,
                 evidence=item_evidence)
    if progress:
        progress.complete_step("select_item")
        progress.begin_step("set_quantity")
    quantity = _set_quantity(form, record.quantity_text)
    journal.emit("field", "OK", "请购数量已写入并读回",
                 step="set_quantity", field="请购数量", value=quantity)
    if progress:
        progress.complete_step("set_quantity")

    item_cell = _grid_cell(form.grid_scope, "品号", 0)
    row_control = item_cell.GetParentControl()
    row_value = _control_value(row_control) or ""
    journal.emit(
        "row_readback", "OK", "表体行关键字段已读回；仓库由 E10 测试账套自动带出",
        step="verify_warehouse", item=record.item_no, quantity=quantity,
        row_serialized_value=row_value,
        warehouse_mode="E10_AUTO_DEFAULT_NOT_IN_XLSX")
    journal.emit(
        "step", "SKIPPED", "输入未提供仓库；保留 E10 自动默认仓库并记录读回",
        step="verify_warehouse", reason="not_in_input_e10_auto_default")
    if progress:
        progress.skip_step("verify_warehouse", "not_in_input_e10_auto_default")
    return {
        "document_values": doc_values,
        "applicant_values": applicant_values,
        "required_date": required_date,
        "correlation_field": "备注",
        "correlation_value": correlation_readback,
        "identity_strength": IdentityStrength.REQUEST_FIELD_MATCH.value,
        "item": record.item_no,
        "quantity": quantity,
        "row_serialized_value": row_value,
    }


def _unique_materialized_named(root, control_type: str, name: str, key: str):
    matches = []
    for control, _depth in _walk_controls(root, max_depth=20, limit=800):
        try:
            if (control.ControlTypeName == control_type
                    and str(control.Name or "") == name
                    and _materialized(control)):
                matches.append(control)
        except Exception:
            continue
    if len(matches) != 1:
        raise RuntimeError(
            f"selector {key!r} 实体化候选期望唯一，实际 {len(matches)} 个")
    return matches[0]


@_timed_wait("unique_materialized_named")
def _wait_unique_materialized_named(
        root, control_type: str, name: str, key: str, *, timeout: float = 15.0):
    """Wait only for zero candidates; ambiguity remains an immediate failure."""
    deadline = time.monotonic() + timeout
    last_error = None
    while time.monotonic() < deadline:
        try:
            return _unique_materialized_named(root, control_type, name, key)
        except RuntimeError as exc:
            last_error = exc
            if "实际 0 个" not in str(exc):
                raise
        time.sleep(0.3)
    raise RuntimeError(f"selector {key!r} 等待实体化超时: {last_error}")


def _select_left_tree_item(main, name: str) -> dict:
    """Materialize and click one semantic left-navigation tree item."""
    import uiautomation as auto
    from e10_query import get_menu_panel

    scope = get_menu_panel(main)

    def resolve():
        common = dict(searchFromControl=scope, searchDepth=30, Name=name)
        first = auto.TreeItemControl(foundIndex=1, **common)
        if not first.Exists(3, 0.2):
            raise RuntimeError(f"模块树里找不到 {name!r}")
        second = auto.TreeItemControl(foundIndex=2, **common)
        if second.Exists(0.3, 0.1):
            raise RuntimeError(f"模块树 {name!r} 命中多个，拒绝选择")
        return first

    item = resolve()
    before_rect = item.BoundingRectangle
    scroll_evidence = None
    method = "TreeItem.Click"
    if not _materialized(item):
        pattern = item.GetScrollItemPattern()
        if pattern is None:
            raise RuntimeError(
                f"模块树 {name!r} 未实体化且不支持 ScrollItemPattern")
        pattern.ScrollIntoView()
        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            item = resolve()
            if _materialized(item):
                break
            time.sleep(0.2)
        if not _materialized(item):
            raise RuntimeError(
                f"模块树 {name!r} ScrollIntoView 后仍未实体化")
        method = "ScrollItemPattern.ScrollIntoView + TreeItem.Click"
        scroll_evidence = {
            "pattern": "ScrollItemPattern",
            "before_rect": (
                before_rect.left, before_rect.top,
                before_rect.right, before_rect.bottom),
            "after_rect": (
                item.BoundingRectangle.left, item.BoundingRectangle.top,
                item.BoundingRectangle.right, item.BoundingRectangle.bottom),
        }
    item.Click()
    return {
        "tree_item": name,
        "method": method,
        "materialized_before_click": True,
        "scroll_evidence": scroll_evidence,
    }


def _scroll_navigation_for_target(root, target_name: str) -> dict:
    """Use the unique horizontal ScrollPattern; never drag by coordinates."""
    import uiautomation as auto

    candidates = []
    for control, _depth in _walk_controls(root, max_depth=12, limit=500):
        try:
            pattern = control.GetPattern(auto.PatternId.ScrollPattern)
            if (pattern is not None and bool(pattern.HorizontallyScrollable)
                    and _materialized(control)):
                candidates.append((control, pattern))
        except Exception:
            continue
    if len(candidates) != 1:
        evidence = [
            {"type": str(control.ControlTypeName or ""),
             "name": str(control.Name or "")[:160],
             "automation_id": str(control.AutomationId or "")}
            for control, _pattern in candidates
        ]
        raise RuntimeError(
            f"导航页水平 ScrollPattern 期望唯一，实际 {len(candidates)} 个: {evidence}")
    control, pattern = candidates[0]
    before_horizontal = float(pattern.HorizontalScrollPercent)
    before_vertical = float(pattern.VerticalScrollPercent)
    pattern.SetScrollPercent(100.0, before_vertical)
    _wait_unique_materialized_named(
        root, "ButtonControl", target_name, f"nav.function.{target_name}",
        timeout=12)
    return {
        "container_type": str(control.ControlTypeName or ""),
        "container_name": str(control.Name or "")[:160],
        "before_horizontal_percent": before_horizontal,
        "after_horizontal_percent": float(pattern.HorizontalScrollPercent),
        "horizontal_view_size": float(pattern.HorizontalViewSize),
        "vertical_percent_preserved": before_vertical,
        "pattern": "ScrollPattern.SetScrollPercent",
        "restore": lambda: pattern.SetScrollPercent(
            before_horizontal, before_vertical),
    }


def _open_new_requisition_from_main(
        journal: RunJournal, progress: ProgressHeartbeat | None = None):
    import uiautomation as auto
    from pywinauto import Desktop
    from e10_query import ensure_ready, select_group

    desktop = Desktop(backend="uia")
    if progress:
        progress.begin_step("ensure_ready")
    main = ensure_ready(desktop)
    e10_pid = int(main.ProcessId)
    journal.emit("step", "OK", "E10 主页面前置检查通过", step="ensure_ready")
    if progress:
        progress.complete_step("ensure_ready")
        progress.begin_step("open_new")
    navigation_restore = None
    select_group(main)
    try:
        function_button = _unique_materialized_named(
            main, "ButtonControl", "维护请购单", "nav.function.维护请购单")
    except RuntimeError as exc:
        if "实际 0 个" not in str(exc):
            raise
        navigation = _select_left_tree_item(main, "采购管理")
        journal.emit(
            "navigation_state", "OK",
            "已选择左侧采购管理并请求渲染右侧采购流程图",
            step="open_new", expected_right_button="维护请购单",
            **navigation)
        try:
            function_button = _wait_unique_materialized_named(
                main, "ButtonControl", "维护请购单",
                "nav.function.维护请购单", timeout=10)
        except RuntimeError as wait_error:
            if "实际 0 个" not in str(wait_error):
                raise
            scroll_evidence = _scroll_navigation_for_target(
                main, "维护请购单")
            navigation_restore = scroll_evidence.pop("restore")
            journal.emit(
                "navigation_scroll", "OK",
                "使用唯一水平 ScrollPattern 使维护请购单按钮物化",
                step="open_new", **scroll_evidence)
            function_button = _unique_materialized_named(
                main, "ButtonControl", "维护请购单",
                "nav.function.维护请购单")
    try:
        function_button.Click()
        browse_window = _wait_top_window(
            "浏览 - 维护请购单", timeout=15, expected_pid=e10_pid,
            journal=journal)
    finally:
        if navigation_restore is not None:
            navigation_restore()
    browse_hwnd = int(browse_window["hwnd"])
    browse = desktop.window(handle=browse_hwnd)
    browse_roots = [browse]
    browse_root = auto.ControlFromHandle(browse_hwnd)
    new_button = _unique_descendant(
        browse_root, auto.ButtonControl, "requisition.browse.new", Name="新建")
    new_button.Click()
    edit_window = _wait_top_window(
        "维护请购单", timeout=15, expected_pid=e10_pid, journal=journal)
    form = attach_live_form(expected_pid=e10_pid)
    if _control_value(form.doc_no_edit):
        raise RuntimeError("新建后的请购单单号应为空，当前已存在值")
    journal.emit(
        "step", "OK", "已从采购管理打开维护请购单并新建空白单",
        step="open_new", browse_hwnd=browse_hwnd,
        edit_hwnd=int(edit_window["hwnd"]))
    if progress:
        progress.complete_step("open_new")
    return desktop, browse, browse_roots, form


def _close_probe_form_without_save(form: LiveForm, initial_handles: set[int],
                                   journal: RunJournal) -> dict:
    """Close only the blank form created by the correlation probe."""
    import uiautomation as auto
    from e10_recorder import enumerate_top_windows

    hwnd = int(form.window["hwnd"])
    pid = int(form.window["pid"])
    _post_close(hwnd)
    clicked_no = False
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline and not _wait_hwnd_closed(
            hwnd, timeout=0.2):
        prompts = []
        for row in enumerate_top_windows():
            if (int(row.get("pid", 0)) != pid
                    or int(row.get("hwnd", 0)) in initial_handles
                    or int(row.get("hwnd", 0)) == hwnd
                    or not _materialized_top_window(row)):
                continue
            root = auto.ControlFromHandle(int(row["hwnd"]))
            no = auto.ButtonControl(
                searchFromControl=root, searchDepth=12,
                foundIndex=1, AutomationId="btnNo")
            if no.Exists(0.2, 0.1) and _materialized(no):
                prompts.append((row, no))
        if len(prompts) > 1:
            raise RuntimeError("关联字段探针关闭时出现多个‘否’确认按钮")
        if len(prompts) == 1:
            prompts[0][1].Click()
            clicked_no = True
        time.sleep(0.2)
    if not _wait_hwnd_closed(hwnd, timeout=1):
        raise RuntimeError("关联字段探针创建的空白表单未能不保存关闭")
    evidence = {
        "edit_hwnd": hwnd,
        "discard_confirmation_no_clicked": clicked_no,
        "saved": False,
    }
    journal.emit(
        "probe_cleanup", "OK", "探针表单已不保存关闭",
        business_data_modified=False, safe_to_retry=True, **evidence)
    return evidence


def probe_correlation_field_live(report_dir: Path) -> dict:
    """Unsaved probe for a writable and queryable requisition correlation field."""
    import uiautomation as auto
    from e10_query import advanced_query, auto_set_edit_value

    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.probe.correlation_field",
        selector_version=selector_version, risk=RiskLevel.REVERSIBLE_WRITE,
        environment={
            "account_set": "FRKTEST",
            "mode": "unsaved_correlation_field_probe",
            "candidate_scope": "selCtrREMARK",
            "candidate_label": "备注",
        })
    initial_handles = _current_top_handles()
    form = None
    cleanup = None
    marker = f"E10RPA-PROBE-{journal.run_id[:8]}"
    try:
        desktop, browse, browse_roots, form = _open_new_requisition_from_main(
            journal)
        scope = _field_scope(form.root, "selCtrREMARK")
        editor = _unique_native_edit(scope, "备注作用域")
        scope_rect = scope.BoundingRectangle
        edit_rect = editor.BoundingRectangle
        original = _control_value(editor) or ""
        selector_evidence = {
            "label": "备注",
            "scope_control_type": str(scope.ControlTypeName or ""),
            "scope_automation_id": str(scope.AutomationId or ""),
            "scope_rect": [scope_rect.left, scope_rect.top,
                           scope_rect.right, scope_rect.bottom],
            "editor_control_type": str(editor.ControlTypeName or ""),
            "editor_class_name": str(editor.ClassName or ""),
            "editor_automation_id": str(editor.AutomationId or ""),
            "editor_automation_id_policy": "diagnostic_only_session_dynamic",
            "editor_native_handle": int(editor.NativeWindowHandle or 0),
            "editor_rect": [edit_rect.left, edit_rect.top,
                            edit_rect.right, edit_rect.bottom],
            "enabled": bool(editor.IsEnabled),
        }
        if not auto_set_edit_value(editor, marker):
            raise RuntimeError("备注候选字段写入后读回失败")
        marker_readback = _control_value(editor) or ""
        if marker_readback != marker:
            raise RuntimeError(
                f"备注候选字段读回不一致: {marker_readback!r}")
        written_shot = _capture_screen(
            report_dir / f"shots_{journal.run_id[:8]}",
            "correlation_field_written_unsaved")
        if not auto_set_edit_value(editor, original):
            raise RuntimeError("备注候选字段无法恢复原值")
        restored = _control_value(editor) or ""
        if restored != original:
            raise RuntimeError(
                f"备注候选字段恢复读回不一致: {restored!r}")
        journal.emit(
            "correlation_field_write_probe", "OK",
            "备注候选字段可写、可读回并已恢复原值",
            selector=selector_evidence,
            marker=marker, marker_readback=marker_readback,
            original_value=original, restored_value=restored,
            screenshot=str(written_shot), persisted=False,
            business_data_modified=False, safe_to_retry=True)

        cleanup = _close_probe_form_without_save(
            form, initial_handles, journal)
        form = None
        outcome = advanced_query(
            browse, desktop, [("备注", marker)], browse_roots,
            shots_dir=report_dir / f"shots_{journal.run_id[:8]}",
            journal=journal)
        query_details = dict(outcome.details or {})
        field_selectable = outcome.code not in {
            "DROPDOWN_ITEM_NOT_FOUND", "SUBMISSION_NOT_ACKNOWLEDGED"}
        if not field_selectable:
            raise RuntimeError(
                f"备注字段未能作为高级查询字段使用: {outcome.code}")
        decision = "ALPHA_CANDIDATE"
        report = {
            "status": "PROBE_COMPLETE",
            "decision": decision,
            "candidate": "selCtrREMARK",
            "label": "备注",
            "writable": True,
            "readback_verified": True,
            "restored": True,
            "persisted": False,
            "advanced_query_field_selectable": True,
            "query_status": outcome.status.value,
            "query_code": outcome.code,
            "query_details": query_details,
            "selector_evidence": selector_evidence,
            "cleanup": cleanup,
            "journal": str(journal.path),
            "limitation": (
                "字段可写且可用于高级查询已证；保存后按关联号精确命中"
                "仍须由唯一一张身份链 E2E 证明"),
        }
        report_path = report_dir / f"correlation_probe_{journal.run_id}.json"
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        report["report"] = str(report_path)
        journal.emit(
            "run_finished", "PROBE_COMPLETE",
            "关联字段探针完成，备注字段进入方案 α 候选",
            business_data_modified=False, safe_to_retry=True,
            decision=decision, report=str(report_path), persisted=False)
        return report
    except Exception as exc:
        if form is not None:
            try:
                cleanup = _close_probe_form_without_save(
                    form, initial_handles, journal)
            except Exception as cleanup_exc:
                journal.emit(
                    "probe_cleanup", "FAILED", str(cleanup_exc),
                    observation_code="PROBE_CLEANUP_FAILED",
                    business_data_modified=False, safe_to_retry=False)
        journal.emit(
            "run_finished", "FAILED", str(exc),
            business_data_modified=False, safe_to_retry=True,
            error_type=type(exc).__name__, cleanup=cleanup)
        raise
    finally:
        # Close only titled windows created after the probe began.  A saved
        # document is impossible because this workflow has no Save actuator.
        _cleanup_created_titled_windows(initial_handles)


def _current_top_handles() -> set[int]:
    from e10_recorder import enumerate_top_windows

    return {int(row["hwnd"]) for row in enumerate_top_windows()}


def _new_titled_windows(before: set[int], *, pid: int) -> list[dict]:
    from e10_recorder import enumerate_top_windows

    result = []
    for row in enumerate_top_windows():
        if (int(row.get("pid", 0)) != int(pid)
                or int(row["hwnd"]) in before
                or not _materialized_top_window(row)
                or not str(row.get("title", "")).strip()):
            continue
        result.append({key: row.get(key) for key in ("hwnd", "pid", "title")})
    return result


def _cleanup_created_titled_windows(initial_handles: set[int]) -> list[dict]:
    """Close only titled E10 windows whose HWND did not exist at run start."""
    from e10_recorder import enumerate_top_windows
    from e10_query import close_gracefully, get_digiwin_pids

    pids = set(get_digiwin_pids())
    results = []
    for row in enumerate_top_windows():
        hwnd = int(row["hwnd"])
        title = str(row.get("title", "")).strip()
        if (hwnd in initial_handles or int(row.get("pid", 0)) not in pids
                or not title or title.startswith("鼎捷ERP E10")):
            continue
        closed = close_gracefully({"hwnd": hwnd}, title)
        results.append({"hwnd": hwnd, "title": title, "closed": closed})
    return results


def execute_create_from_main(record: RequisitionRecord, report_dir: Path,
                             input_path: Path | None = None, *,
                             request_no: str,
                             dataset_epoch: str,
                             request_registry: RequestRegistry,
                             gate_decision=None) -> dict:
    """End-to-end FRKTEST create with one Save and UI test-database reconcile."""
    if (request_registry is None or gate_decision is None
            or gate_decision.action != GateAction.ALLOW_WRITE):
        raise RuntimeError("防重闸未授权本次写入")
    import uiautomation as auto
    from e10_query import advanced_query, close_gracefully, scrape_grid

    dataset_epoch = require_dataset_epoch(dataset_epoch)
    correlation = correlation_value(request_no)
    document_identity_key = identity_key(dataset_epoch, request_no)

    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    _print_plan("e10.requisition.create.end_to_end",
                warehouse_provided=record.warehouse is not None)
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.create.end_to_end",
        selector_version=selector_version, risk=RiskLevel.COMMIT,
        environment={
            "account_set": "FRKTEST",
            "mode": "test_database_create",
            "reconcile_channel": "E10_UI_TEST_DATABASE",
            "middledb_available_for_account_set": False,
            "input_file": str(input_path.resolve()) if input_path else None,
            "input_hash": _input_sha256(input_path),
            "request_no": request_no,
            "dataset_epoch": dataset_epoch,
            "document_identity_key": document_identity_key,
            "identity_strength": IdentityStrength.REQUEST_FIELD_MATCH.value,
            "correlation_field": "备注",
            "correlation_value": correlation,
            "payload_fingerprint": record.payload_fingerprint,
        },
    )
    progress = ProgressHeartbeat(
        interval=10.0, output=lambda text: print(text, flush=True),
        event_sink=_step_timing_sink(journal))
    progress.start()
    save_issued = False
    pending_registered = False
    try:
        _require_supported_warehouse(record)
        initial_handles = _current_top_handles()
        desktop, browse, browse_roots, form = _open_new_requisition_from_main(
            journal, progress)
        filled = _fill_current_form(
            form, record, journal, progress, correlation=correlation)
        filled_shot = _capture_screen(
            report_dir / f"shots_{journal.run_id[:8]}", "filled_before_validation")
        journal.emit("screenshot", "OK", "保存前字段截图", path=str(filled_shot))

        progress.begin_step("validate")
        validation = validate_filled_form(
            record, report_dir, correlation=correlation)
        journal.emit("validation", "OK", "E10 普通校验已单击，关键字段保持一致",
                     step="validate",
                     validation_status=validation["status"],
                     doc_no=validation["doc_no"],
                     item=validation["item"],
                     quantity=validation["quantity"],
                     correlation_field="备注",
                     correlation_value=validation["correlation_value"],
                     identity_strength=IdentityStrength.REQUEST_FIELD_MATCH.value,
                     screenshot=validation["screenshot"],
                     unexpected_windows=validation["unexpected_windows"])
        progress.complete_step("validate")
        doc_no = _control_value(form.doc_no_edit) or ""
        if not doc_no:
            raise RuntimeError("保存前 E10 未预分配单号")

        context = {
            "form": form,
            "record": record,
            "doc_no": doc_no,
            "desktop": desktop,
            "browse": browse,
            "browse_roots": browse_roots,
            "journal": journal,
            "report_dir": report_dir,
            "correlation_value": correlation,
        }

        def save_precondition(ctx):
            item = _control_value(_grid_cell(ctx["form"].grid_scope, "品号", 0))
            quantity = (_control_value(_grid_cell(
                ctx["form"].grid_scope, "请购数量", 0)) or "").replace(",", "")
            if item != ctx["record"].item_no or quantity != ctx["record"].quantity_text:
                raise RuntimeError("保存前关键字段最终读回不一致")
            correlation_actual = _control_value(_unique_native_edit(
                _field_scope(ctx["form"].root, "selCtrREMARK"),
                "备注作用域")) or ""
            if correlation_actual != ctx["correlation_value"]:
                raise RuntimeError("保存前备注关联号最终读回不一致")

        def locate_save(ctx):
            return _unique_descendant(
                ctx["form"].root, auto.ButtonControl,
                "requisition.save", Name="保存")

        def act_save(ctx, button):
            nonlocal save_issued
            ctx["before_save_windows"] = _current_top_handles()
            button.Click()
            save_issued = True

        def readback_save(ctx, _button):
            time.sleep(2)
            shot = _capture_screen(
                report_dir / f"shots_{journal.run_id[:8]}", "after_save_click")
            return {
                "screenshot": str(shot),
                "unexpected": _new_titled_windows(
                    ctx["before_save_windows"], pid=int(ctx["form"].window["pid"])),
            }

        def expect_save(_ctx, readback):
            if readback["unexpected"]:
                raise RuntimeError(
                    f"保存后出现新提示窗口: {readback['unexpected']}")

        def reconcile_save(ctx):
            form_hwnd = int(ctx["form"].window["hwnd"])
            _post_close(form_hwnd)
            if not _wait_hwnd_closed(form_hwnd, timeout=8):
                return ReconcileEvidence(
                    ReconcileStatus.UNKNOWN,
                    source="E10_UI_TEST_DATABASE",
                    external=False,
                    details={"reason": "保存后关闭编辑窗出现确认或未关闭",
                             "doc_no": ctx["doc_no"]})
            try:
                outcome = advanced_query(
                    ctx["browse"], ctx["desktop"],
                    [("备注", ctx["correlation_value"])],
                    ctx["browse_roots"],
                    shots_dir=report_dir / f"shots_{journal.run_id[:8]}",
                    journal=journal,
                    result_expected={"单号": ctx["doc_no"]})
            except Exception as exc:
                return ReconcileEvidence(
                    ReconcileStatus.UNKNOWN,
                    source="E10_UI_TEST_DATABASE",
                    external=False,
                    details={"reason": f"UI 查询核对失败: {exc}",
                             "doc_no": ctx["doc_no"]})
            details = {
                "doc_no": ctx["doc_no"],
                "dataset_epoch": dataset_epoch,
                "request_no": request_no,
                "document_identity_key": document_identity_key,
                "identity_strength": IdentityStrength.REQUEST_FIELD_MATCH.value,
                "correlation_field": "备注",
                "correlation_value": ctx["correlation_value"],
                "query_status": outcome.status.value,
                "query_code": outcome.code,
                "query_message": outcome.message,
                "account_set": "FRKTEST",
                "middledb_skipped": "测试数据库不进入 MIDDLEDB",
            }
            if outcome.status == QueryStatus.FOUND:
                if len(outcome.matched_rows) != 1:
                    details["reason"] = (
                        "备注关联号查询未得到唯一匹配: "
                        f"{len(outcome.matched_rows)}")
                    return ReconcileEvidence(
                        ReconcileStatus.UNKNOWN,
                        source="E10_UI_TEST_DATABASE",
                        external=False, details=details)
                try:
                    columns, table, _snapshot = scrape_grid(ctx["browse"])
                    full = GridSnapshot.from_table(columns, table)
                    doc_rows = full.matching_rows({"单号": ctx["doc_no"]})
                except Exception as exc:
                    details["reason"] = f"关联号命中后读取完整行失败: {exc}"
                    return ReconcileEvidence(
                        ReconcileStatus.UNKNOWN,
                        source="E10_UI_TEST_DATABASE",
                        external=False, details=details)
                if len(doc_rows) != 1:
                    details["reason"] = (
                        "关联号命中，但目标单号在结果中不唯一: "
                        f"{len(doc_rows)}")
                    details["result_columns"] = columns
                    return ReconcileEvidence(
                        ReconcileStatus.UNKNOWN,
                        source="E10_UI_TEST_DATABASE",
                        external=False, details=details)
                details["matched_business_row"] = dict(doc_rows[0])
                shot = _capture_screen(
                    report_dir / f"shots_{journal.run_id[:8]}", "browser_reconciled")
                details["screenshot"] = str(shot)
                browse_hwnd = int(ctx["browse"].element_info.handle)
                details["browse_closed"] = close_gracefully(
                    {"hwnd": browse_hwnd}, "浏览 - 维护请购单")
                return ReconcileEvidence(
                    ReconcileStatus.CONFIRMED,
                    source="E10_UI_TEST_DATABASE",
                    external=False,
                    details=details)
            if outcome.status == QueryStatus.EMPTY:
                return ReconcileEvidence(
                    ReconcileStatus.NOT_APPLIED,
                    source="E10_UI_TEST_DATABASE",
                    external=False,
                    details=details)
            return ReconcileEvidence(
                ReconcileStatus.UNKNOWN,
                source="E10_UI_TEST_DATABASE",
                external=False,
                details=details)

        request_registry.record_pending(
            gate_decision, record.payload_fingerprint,
            run_id=journal.run_id, doc_no=doc_no,
            dataset_epoch=dataset_epoch,
            evidence_jsonl=str(journal.path.resolve()))
        pending_registered = True
        journal.emit(
            "request_gate", "PENDING", "业务请求已先登记，允许单次保存",
            request_no=request_no, attempt=gate_decision.attempt,
            dataset_epoch=dataset_epoch,
            document_identity_key=document_identity_key,
            identity_strength=IdentityStrength.REQUEST_FIELD_MATCH.value,
            correlation_field="备注", correlation_value=correlation,
            payload_fingerprint=record.payload_fingerprint)

        intent = WriteIntent(
            request_no=request_no,
            payload_fingerprint=record.payload_fingerprint,
            document_key=doc_no,
            action="SAVE_REQUISITION_FRKTEST",
            before_values={"persisted": False, "doc_no_preallocated": doc_no},
            payload={**record.as_business_mapping(), "备注": correlation},
        )
        step = StepSpec(
            name="save_requisition_once",
            risk=RiskLevel.COMMIT,
            precondition=save_precondition,
            locate=locate_save,
            act=act_save,
            readback=readback_save,
            expect=expect_save,
            reconcile=reconcile_save,
            writes_business_data=True,
            intent=intent,
            human_confirmed=True,
            require_external_reconcile=False,
        )
        progress.begin_step("save_requisition_once")
        step_result = StepExecutor(journal).execute(step, context)
        if step_result.status == StepStatus.SUCCEEDED:
            progress.complete_step("save_requisition_once")
        else:
            progress.fail_step(RuntimeError(step_result.message))
        cleanup = []
        if step_result.status == StepStatus.SUCCEEDED:
            progress.begin_step("cleanup")
            cleanup = _cleanup_created_titled_windows(initial_handles)
            journal.emit(
                "cleanup", "OK" if all(row["closed"] for row in cleanup) else "FAILED",
                "已按运行前 HWND 快照清理本次新建的剩余 E10 窗口",
                step="cleanup", created_windows=cleanup)
            if all(row["closed"] for row in cleanup):
                progress.complete_step("cleanup")
            else:
                progress.fail_step(RuntimeError("owned window cleanup failed"))
        status = (
            "SAVED_CONFIRMED"
            if step_result.status == StepStatus.SUCCEEDED
            else "NOT_APPLIED"
            if step_result.reconcile_status == ReconcileStatus.NOT_APPLIED
            else "UNKNOWN")
        if pending_registered:
            registry_status = {
                ReconcileStatus.CONFIRMED: RequestStatus.CONFIRMED,
                ReconcileStatus.NOT_APPLIED: RequestStatus.NOT_APPLIED,
                ReconcileStatus.UNKNOWN: RequestStatus.UNKNOWN,
                None: RequestStatus.UNKNOWN,
            }[step_result.reconcile_status]
            request_registry.append(
                request_no=request_no,
                campaign_id=gate_decision.campaign_id,
                dataset_epoch=dataset_epoch,
                payload_fingerprint=record.payload_fingerprint,
                status=registry_status,
                run_id=journal.run_id,
                attempt=gate_decision.attempt,
                doc_no=doc_no,
                evidence_jsonl=str(journal.path.resolve()),
                details={"reconcile_status": (
                    step_result.reconcile_status.value
                    if step_result.reconcile_status else None)},
            )
        journal.emit(
            "run_finished", status, step_result.message,
            business_data_modified=step_result.business_data_modified,
            safe_to_retry=step_result.safe_to_retry,
            doc_no=doc_no, reconcile_status=(
                step_result.reconcile_status.value
                if step_result.reconcile_status else None),
            failure_code=(
                FailureCode.WRITE_NOT_APPLIED.value
                if step_result.reconcile_status == ReconcileStatus.NOT_APPLIED
                else None),
            dataset_epoch=dataset_epoch,
            request_no=request_no,
            document_identity_key=document_identity_key,
            identity_strength=IdentityStrength.REQUEST_FIELD_MATCH.value,
            correlation_field="备注", correlation_value=correlation,
            record=record.as_business_mapping())
        return {
            "status": status,
            "doc_no": doc_no,
            "dataset_epoch": dataset_epoch,
            "request_no": request_no,
            "document_identity_key": document_identity_key,
            "identity_strength": IdentityStrength.REQUEST_FIELD_MATCH.value,
            "correlation_field": "备注",
            "correlation_value": correlation,
            "journal": str(journal.path),
            "filled_screenshot": str(filled_shot),
            "reconcile_status": (step_result.reconcile_status.value
                                 if step_result.reconcile_status else None),
            "reconcile_evidence": dict(step_result.evidence),
            "business_data_modified": step_result.business_data_modified,
            "safe_to_retry": step_result.safe_to_retry,
            "cleanup": cleanup,
            "filled": filled,
        }
    except Exception as exc:
        progress.fail_step(exc)
        if pending_registered:
            request_registry.append(
                request_no=request_no,
                campaign_id=gate_decision.campaign_id,
                dataset_epoch=dataset_epoch,
                payload_fingerprint=record.payload_fingerprint,
                status=(RequestStatus.UNKNOWN if save_issued
                        else RequestStatus.NOT_APPLIED),
                run_id=journal.run_id,
                attempt=gate_decision.attempt,
                doc_no=locals().get("doc_no"),
                evidence_jsonl=str(journal.path.resolve()),
                details={"error": str(exc), "write_request_issued": save_issued},
            )
        journal.emit(
            "run_finished", "FAILED", str(exc),
            business_data_modified=False,
            safe_to_retry=not save_issued,
            write_request_issued=save_issued,
            error_type=type(exc).__name__)
        raise
    finally:
        progress.stop()


class _LiveVerifyActor:
    """UI adapter for the read-only verify protocol."""

    def __init__(self, journal: RunJournal, report_dir: Path,
                 progress: ProgressHeartbeat, *, request_no: str | None = None):
        self.journal = journal
        self.report_dir = report_dir
        self.progress = progress
        self.initial_handles = _current_top_handles()
        self.desktop = None
        self.browse = None
        self.browse_roots = []
        self.navigation_restore = None
        self.e10_pid = None
        self.request_no = request_no

    def open(self) -> None:
        from pywinauto import Desktop
        from e10_query import ensure_ready, select_group

        self.desktop = Desktop(backend="uia")
        self.progress.begin_step("ensure_ready")
        try:
            main = ensure_ready(self.desktop)
            self.e10_pid = int(main.ProcessId)
        except Exception as exc:
            self.progress.fail_step(exc)
            raise
        self.journal.emit(
            "step", "OK", "E10 主页面前置检查通过", step="ensure_ready")
        self.progress.complete_step("ensure_ready")

        self.progress.begin_step("open_browse")
        try:
            select_group(main)
            try:
                function_button = _unique_materialized_named(
                    main, "ButtonControl", "维护请购单", "nav.function.维护请购单")
            except RuntimeError as exc:
                if "实际 0 个" not in str(exc):
                    raise
                # Clicking the left navigation row selects the module and
                # causes E10 to render its clickable process diagram on the
                # right.  ExpandCollapse only opens children and does not
                # perform that selection.
                navigation = _select_left_tree_item(main, "采购管理")
                self.journal.emit(
                    "navigation_state", "OK",
                    "已点击左侧采购管理并请求渲染右侧采购流程图",
                    step="open_browse", expected_right_button="维护请购单",
                    **navigation)
                try:
                    function_button = _wait_unique_materialized_named(
                        main, "ButtonControl", "维护请购单",
                        "nav.function.维护请购单", timeout=10)
                except RuntimeError as wait_error:
                    if "实际 0 个" not in str(wait_error):
                        raise
                    scroll_evidence = _scroll_navigation_for_target(
                        main, "维护请购单")
                    self.navigation_restore = scroll_evidence.pop("restore")
                    self.journal.emit(
                        "navigation_scroll", "OK",
                        "使用唯一水平 ScrollPattern 使维护请购单按钮物化",
                        step="open_browse", **scroll_evidence)
                    function_button = _unique_materialized_named(
                        main, "ButtonControl", "维护请购单",
                        "nav.function.维护请购单")
            function_button.Click()
            browse_window = _wait_top_window(
                "浏览 - 维护请购单", timeout=15,
                expected_pid=self.e10_pid, journal=self.journal)
            self.browse = self.desktop.window(handle=int(browse_window["hwnd"]))
            if self.navigation_restore is not None:
                self.navigation_restore()
                self.navigation_restore = None
        except Exception as exc:
            self.progress.fail_step(exc)
            raise
        self.browse_roots.append(self.browse)
        self.browse.set_focus()
        self.browse_roots[-1] = self.browse
        self.journal.emit(
            "step", "OK", "已打开维护请购单浏览窗口",
            step="open_browse",
            browse_hwnd=int(self.browse.element_info.handle))
        self.progress.complete_step("open_browse")

    def query(self, document_no: str) -> QueryOutcome:
        from e10_query import advanced_query, scrape_grid

        query_field = "备注" if self.request_no else "单号"
        query_value = self.request_no or document_no
        # Keep the verified workflow step name stable; query_field records
        # whether this run used the strong request correlation or bare doc no.
        step_name = "query_by_document_no"
        self.progress.begin_step(step_name)
        try:
            outcome = advanced_query(
                self.browse, self.desktop, [(query_field, query_value)],
                self.browse_roots,
                shots_dir=self.report_dir / f"shots_{self.journal.run_id[:8]}",
                journal=self.journal,
                result_expected={"单号": document_no},
            )
            if outcome.status == QueryStatus.FOUND:
                columns, table, _snapshot = scrape_grid(self.browse)
                full = GridSnapshot.from_table(columns, table)
                matches = full.matching_rows({"单号": document_no})
                if len(matches) != 1:
                    outcome = QueryOutcome(
                        QueryStatus.FAILED,
                        "IDENTITY_NOT_UNIQUE",
                        "查询结果未唯一命中预期单号",
                        matched_rows=(),
                        safe_to_retry=True,
                        details={
                            **dict(outcome.details or {}),
                            "expected_doc_no": document_no,
                            "matched_doc_rows": len(matches),
                            "query_field": query_field,
                            "query_value": query_value,
                        },
                    )
                else:
                    outcome = QueryOutcome(
                        outcome.status, outcome.code, outcome.message,
                        matched_rows=outcome.matched_rows,
                        safe_to_retry=outcome.safe_to_retry,
                        details={
                            **dict(outcome.details or {}),
                            "query_field": query_field,
                            "query_value": query_value,
                            "matched_business_row": dict(matches[0]),
                        },
                    )
        except Exception as exc:
            self.progress.fail_step(exc)
            raise
        details = dict(outcome.details or {})
        shot = _capture_screen(
            self.report_dir / f"shots_{self.journal.run_id[:8]}",
            "verify_result")
        self.journal.emit(
            "query_result", outcome.status.value, outcome.message,
            step=step_name, doc_no=document_no,
            request_no=self.request_no,
            query_field=query_field, query_value=query_value,
            query_code=outcome.code,
            reported_row_count=details.get("reported_row_count"),
            status_texts=details.get("status_texts"),
            matched_status_texts=details.get("matched_status_texts"),
            screenshot=str(shot),
            business_data_modified=False,
            safe_to_retry=outcome.status != QueryStatus.FAILED,
        )
        self.progress.complete_step(step_name)
        return outcome

    def cleanup(self) -> None:
        self.progress.begin_step("cleanup")
        if self.navigation_restore is not None:
            self.navigation_restore()
            self.navigation_restore = None
        cleanup = _cleanup_created_titled_windows(self.initial_handles)
        okay = all(row["closed"] for row in cleanup)
        self.journal.emit(
            "cleanup", "OK" if okay else "FAILED",
            "只清理本次 verify 新建且仍存活的 E10 窗口",
            step="cleanup", created_windows=cleanup)
        if not okay:
            error = RuntimeError("verify owned window cleanup failed")
            self.progress.fail_step(error)
            raise error
        self.progress.complete_step("cleanup")


def verify_requisition_live(document_no: str, report_dir: Path, *,
                            request_no: str | None = None,
                            payload_fingerprint: str | None = None,
                            dataset_epoch: str | None = None) -> dict:
    """Read-only exact query mapped to the existing reconciliation vocabulary."""
    if bool(request_no) != bool(dataset_epoch):
        raise ValueError(
            "身份核对必须同时提供 request_no 与 dataset_epoch，或两者都不提供")
    strong_identity = bool(request_no and dataset_epoch)
    if strong_identity:
        dataset_epoch = require_dataset_epoch(dataset_epoch)
        request_no = correlation_value(request_no)
        document_identity_key = identity_key(dataset_epoch, request_no)
        identity_strength = IdentityStrength.REQUEST_FIELD_MATCH
    else:
        document_identity_key = None
        identity_strength = IdentityStrength.DOC_NO_ONLY
    selector_version = json.loads(
        (ROOT / "selectors.json").read_text(encoding="utf-8"))["_meta"]["selector_version"]
    _print_plan("e10.requisition.verify")
    journal = _make_journal(
        report_dir, workflow_id="e10.requisition.verify",
        selector_version=selector_version, risk=RiskLevel.READ_ONLY,
        environment={
            "account_set": "FRKTEST",
            "mode": "read_only_verify",
            "reconcile_channel": "E10_UI_TEST_DATABASE",
            "external": False,
            "request_no": request_no,
            "dataset_epoch": dataset_epoch,
            "document_identity_key": document_identity_key,
            "identity_strength": identity_strength.value,
            "correlation_field": "备注" if strong_identity else None,
            "payload_fingerprint": payload_fingerprint,
        },
    )
    progress = ProgressHeartbeat(
        interval=10.0, output=lambda text: print(text, flush=True),
        event_sink=_step_timing_sink(journal))
    progress.start()
    try:
        actor = _LiveVerifyActor(
            journal, report_dir, progress, request_no=request_no)
        evidence = execute_verify_protocol(actor, document_no)
        evidence = enforce_identity_strength(evidence, identity_strength)
        safe = evidence.status != ReconcileStatus.UNKNOWN
        journal.emit(
            "reconcile",
            "OK" if evidence.status == ReconcileStatus.CONFIRMED else
            "EMPTY" if evidence.status == ReconcileStatus.NOT_APPLIED else "FAILED",
            "独立只读核对完成",
            reconcile_status=evidence.status.value,
            reconcile_source=evidence.source,
            external=evidence.external,
            reconcile_details=dict(evidence.details),
            dataset_epoch=dataset_epoch, request_no=request_no,
            document_identity_key=document_identity_key,
            identity_strength=identity_strength.value,
            business_data_modified=False,
            safe_to_retry=safe,
        )
        journal.emit(
            "run_finished", evidence.status.value, "独立只读核对结束",
            business_data_modified=False, safe_to_retry=safe,
            doc_no=document_no, reconcile_status=evidence.status.value,
            reconcile_source=evidence.source, external=evidence.external,
            dataset_epoch=dataset_epoch, request_no=request_no,
            document_identity_key=document_identity_key,
            identity_strength=identity_strength.value)
        return {
            "status": evidence.status.value,
            "doc_no": document_no,
            "dataset_epoch": dataset_epoch,
            "request_no": request_no,
            "document_identity_key": document_identity_key,
            "identity_strength": identity_strength.value,
            "journal": str(journal.path),
            "reconcile_source": evidence.source,
            "external": evidence.external,
            "details": dict(evidence.details),
            "business_data_modified": False,
            "safe_to_retry": safe,
        }
    except Exception as exc:
        journal.emit(
            "run_finished", "UNKNOWN", str(exc),
            business_data_modified=False, safe_to_retry=True,
            doc_no=document_no, reconcile_status=ReconcileStatus.UNKNOWN.value,
            reconcile_source="E10_UI_TEST_DATABASE", external=False,
            dataset_epoch=dataset_epoch, request_no=request_no,
            document_identity_key=document_identity_key,
            identity_strength=identity_strength.value,
            error_type=type(exc).__name__)
        raise
    finally:
        progress.stop()


def _progress(message: str) -> None:
    print(f"[请购探针] {message}", flush=True)


def _inspect_open_form_worker(output: Path) -> Path:
    """Worker body; the parent process enforces the wall-clock deadline."""
    import uiautomation as auto
    from e10_recorder import enumerate_top_windows

    _progress("1/4 正在枚举 Win32 顶层窗口（只读，不点击、不输入）……")
    matches = [row for row in enumerate_top_windows()
               if row.get("title", "").strip() == "维护请购单"]
    if len(matches) != 1:
        raise RuntimeError(
            f"探针要求恰好一个已打开的‘维护请购单’窗口，实际 {len(matches)} 个")

    window = matches[0]
    hwnd = int(window["hwnd"])
    _progress(f"已找到窗口 hwnd={hwnd}；2/4 正在连接该窗口的 UIA 根节点……")
    root = auto.ControlFromHandle(hwnd)
    if not root:
        raise RuntimeError(f"无法从 hwnd={hwnd} 建立 UIA 根节点")

    _progress(
        f"3/4 正在限深度读取控件树（最大 {PROBE_MAX_DEPTH} 层、"
        f"{PROBE_NODE_LIMIT} 节点）……")
    controls = []
    total_nodes = 0
    ancestor_stack = []
    walk = auto.WalkTree(
        root,
        getChildren=lambda control: control.GetChildren(),
        includeTop=False,
        maxDepth=PROBE_MAX_DEPTH,
    )
    for control, depth, _remaining in walk:
        total_nodes += 1
        if total_nodes > PROBE_NODE_LIMIT:
            raise RuntimeError(
                f"控件树超过安全上限 {PROBE_NODE_LIMIT}，拒绝截断后误判")
        index = total_nodes
        ancestor_stack = ancestor_stack[:max(0, depth - 1)]
        try:
            identity = {
                "control_type": str(control.ControlTypeName or ""),
                "name": str(control.Name or ""),
                "automation_id": str(control.AutomationId or ""),
            }
        except Exception:
            identity = {"control_type": "", "name": "", "automation_id": ""}
        semantic_ancestors = [
            item for item in ancestor_stack
            if item.get("name") or (
                item.get("automation_id")
                and not str(item.get("automation_id")).isdigit())
        ]
        parent = ancestor_stack[-1] if ancestor_stack else None
        ancestor_stack.append(identity)
        if index % 100 == 0:
            _progress(f"已处理 {index} 个节点")
        try:
            rect = control.BoundingRectangle
            if not rect or rect.right <= rect.left or rect.bottom <= rect.top:
                continue
            automation_id = str(control.AutomationId or "")
            controls.append({
                "depth": depth,
                "control_type": str(control.ControlTypeName or ""),
                "name": str(control.Name or ""),
                "automation_id": automation_id,
                "automation_id_stability": _stability(automation_id),
                "class_name": str(control.ClassName or ""),
                "rect": [rect.left, rect.top, rect.right, rect.bottom],
                "enabled": bool(control.IsEnabled),
                "parent": parent,
                "semantic_ancestors": semantic_ancestors,
            })
        except Exception as exc:
            controls.append({"probe_error": f"{type(exc).__name__}: {exc}"})

    payload = {
        "schema_version": 1,
        "timestamp": datetime.now().astimezone().isoformat(timespec="milliseconds"),
        "mode": "read_only_one_shot",
        "window": {"title": window["title"], "handle": hwnd,
                   "pid": window.get("pid")},
        "tree_max_depth": PROBE_MAX_DEPTH,
        "tree_node_limit": PROBE_NODE_LIMIT,
        "node_count": total_nodes,
        "materialized_controls": controls,
        "policy": {
            "numeric_automation_ids_allowed": False,
            "coordinates_allowed_for_actions": False,
            "candidate_requires_restart_verification": True,
        },
    }
    _progress("4/4 正在写入探针证据……")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".partial")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(output)
    _progress(f"完成：{output}")
    return output


def inspect_open_form(report_dir: Path, *, timeout_seconds: float = 45.0) -> Path:
    """Run the UIA snapshot in a killable child and print liveness updates."""
    if timeout_seconds < 5:
        raise ValueError("--probe-timeout 不能小于 5 秒")
    report_dir.mkdir(parents=True, exist_ok=True)
    output = report_dir / f"requisition_form_probe_{datetime.now():%Y%m%d_%H%M%S}.json"
    command = [
        sys.executable,
        "-u",
        str(Path(__file__).resolve()),
        "--_probe-worker",
        str(output),
    ]
    environment = dict(os.environ)
    environment["PYTHONUNBUFFERED"] = "1"
    _progress(
        f"开始一次性扫描；硬超时 {timeout_seconds:g} 秒。期间会持续显示累计耗时。")
    process = subprocess.Popen(command, env=environment)
    started = time.monotonic()
    next_notice = started + 3.0
    try:
        while True:
            return_code = process.poll()
            if return_code is not None:
                if return_code != 0:
                    raise RuntimeError(f"探针子进程失败，退出码 {return_code}")
                if not output.is_file():
                    raise RuntimeError("探针子进程已结束，但没有生成证据文件")
                return output
            now = time.monotonic()
            elapsed = now - started
            if elapsed >= timeout_seconds:
                process.terminate()
                try:
                    process.wait(timeout=3)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait(timeout=3)
                partial = output.with_suffix(output.suffix + ".partial")
                if partial.exists():
                    partial.unlink()
                raise TimeoutError(
                    f"UIA 扫描超过 {timeout_seconds:g} 秒，已终止子进程；"
                    "E10 控件树没有在时限内响应")
            if now >= next_notice:
                _progress(f"仍在工作，累计 {elapsed:.0f} 秒……")
                next_notice = now + 3.0
            time.sleep(0.2)
    except KeyboardInterrupt:
        _progress("收到 Ctrl+C，正在终止本次探针……")
        process.terminate()
        try:
            process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            process.kill()
        raise


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="E10 请购单写入模板：输入校验、录制对齐与只读控件探针")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--sheet", help="工作表名；不填时要求恰好一个可见工作表")
    parser.add_argument("--recording", type=Path,
                        help="用于对齐的录制 JSONL；默认取 runs/recordings 中最新文件")
    parser.add_argument("--validate-only", action="store_true",
                        help="只校验并打印 XLSX 数据")
    parser.add_argument("--inspect-open-form", action="store_true",
                        help="只读扫描一个已打开的‘维护请购单’窗口，不点击不输入")
    parser.add_argument("--probe-correlation-field-live", action="store_true",
                        help="FRKTEST 无保存探测备注字段可写、读回及高级查询可用性")
    parser.add_argument("--preflight-live", action="store_true",
                        help="只读验证当前空白新单的写入前置条件")
    parser.add_argument("--fill-live", action="store_true",
                        help="从 XLSX 填写当前空白新单并读回，但不点击保存")
    parser.add_argument("--resume-fill-live", action="store_true",
                        help="续跑本轮表头已填、表体未填的新单；仍不保存")
    parser.add_argument("--resume-open-item-dialog-live", action="store_true",
                        help="从已打开的普通品号查询窗续跑；仍不保存")
    parser.add_argument("--resume-quantity-live", action="store_true",
                        help="从已精确回填品号的 row 0 续填数量；仍不保存")
    parser.add_argument("--validate-filled-live", action="store_true",
                        help="校验当前精确匹配 XLSX 的已填新单；不保存")
    parser.add_argument("--discard-current-test-draft", action="store_true",
                        help="丢弃与 XLSX 精确匹配的当前 FRKTEST 草稿并回到主页面")
    parser.add_argument("--execute-create-live", action="store_true",
                        help="从 E10 主页面端到端新建、保存、UI 查询核对并退出")
    parser.add_argument("--verify-live", action="store_true",
                        help="只读打开维护请购单并按 --doc-no 独立核对")
    parser.add_argument("--doc-no", help="--verify-live 使用的请购单号")
    parser.add_argument("--request-no",
                        help="写流程业务请求号；可由 XLSX 的同名列替代")
    parser.add_argument("--dataset-epoch",
                        help="FRKTEST 数据集/重置批次标识；写入与强身份核对必填")
    parser.add_argument("--retry-not-applied", action="store_true",
                        help="人工明确允许重跑同请求号的 NOT_APPLIED attempt")
    parser.add_argument("--request-registry", type=Path,
                        default=REQUEST_REGISTRY)
    parser.add_argument("--campaign-id", default="",
                        help="稳定性 campaign 标识；默认空表示非 campaign")
    parser.add_argument("--attest-no-intervention", action="store_true",
                        help="声明本次运行期间不人工操作鼠标键盘")
    parser.add_argument("--desktop-check", action="store_true",
                        help="只读证明当前进程与 E10 是否位于同一交互桌面")
    parser.add_argument("--allow-start-e10", action="store_true",
                        help="桌面已证明可交互时，显式允许原有自动启动 E10 路径")
    parser.add_argument("--allow-readonly-lock-degradation", action="store_true",
                        help="仅 READ_ONLY 流程显式允许 Global 拒绝时降级到 Local 锁")
    parser.add_argument("--scenario-id", choices=[f"S{i}" for i in range(1, 13)],
                        help="campaign 场景编号 S1..S12")
    parser.add_argument("--expected-result",
                        choices=("SUCCESS", "SAFE_REJECT", "OBSERVE"),
                        help="运行前声明该场景期望结果")
    parser.add_argument("--build-test-ledger", action="store_true",
                        help="从 runs/**/*.jsonl 派生 FRKTEST 测试单据台账，不接触 E10")
    parser.add_argument("--scaffold-manual-registration", action="store_true",
                        help="按稳定身份键补齐人工登记空槽；绝不覆盖已有人工字段")
    parser.add_argument("--ledger-output", type=Path,
                        default=ROOT / "runs" / "requisition_test_ledger.json")
    parser.add_argument("--manual-registration", type=Path,
                        default=MANUAL_REGISTRATIONS)
    parser.add_argument("--build-campaign-report", action="store_true",
                        help="从 JSONL 离线生成指定 campaign 的分区报告")
    parser.add_argument("--campaign-report-output", type=Path,
                        default=ROOT / "runs" / "requisition_campaign_report.json")
    parser.add_argument("--dismiss-unexpected-item-dialog", action="store_true",
                        help="仅取消已知的错误品号查询窗口，不修改或保存单据")
    parser.add_argument("--account-set", help="写流程必须显式传 FRKTEST")
    parser.add_argument("--human-present", action="store_true",
                        help="确认测试账套写入时有人在场")
    parser.add_argument("--probe-timeout", type=float, default=45.0,
                        help="只读探针硬超时秒数（默认 45，最小 5）")
    parser.add_argument("--report-dir", type=Path,
                        default=ROOT / "runs" / "requisition_probes")
    parser.add_argument("--_probe-worker", type=Path, help=argparse.SUPPRESS)
    return parser


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    _mutex = None
    code_hash, _code_manifest = compute_code_hash(ROOT)
    _RUNTIME_METADATA.update({
        "campaign_id": args.campaign_id,
        "code_hash": code_hash,
        "attest_no_intervention": bool(args.attest_no_intervention),
        "scenario_id": args.scenario_id or "",
        "expected_result": args.expected_result or "",
        "desktop_guard": None,
        "lock_namespace": None,
        "mutex_denied_reason": None,
    })
    try:
        if args._probe_worker:
            _inspect_open_form_worker(args._probe_worker)
            return 0
        if args.desktop_check:
            from e10_desktop_guard import (DesktopGuardError,
                                           require_interactive_e10)
            from e10_query import get_digiwin_pids
            try:
                decision = require_interactive_e10(
                    get_digiwin_pids(),
                    allow_start_e10=args.allow_start_e10)
                _RUNTIME_METADATA["desktop_guard"] = {
                    "code": decision.code,
                    "message": decision.message,
                    "details": decision.details,
                }
                evidence = _record_desktop_guard_result(
                    args.report_dir, status="READY", code=decision.code,
                    message=decision.message, details=dict(decision.details))
                print(json.dumps({
                    "status": "READY",
                    "code": decision.code,
                    "message": decision.message,
                    "desktop_context": decision.details,
                    "evidence_jsonl": str(evidence),
                }, ensure_ascii=False, indent=2), flush=True)
                return 0
            except DesktopGuardError as exc:
                _RUNTIME_METADATA["desktop_guard"] = {
                    "code": exc.failure_code,
                    "message": str(exc),
                    "details": exc.details,
                }
                evidence = _record_desktop_guard_terminal(
                    args.report_dir, exc)
                print(json.dumps({
                    "status": "BLOCKED",
                    "failure_code": exc.failure_code,
                    "message": str(exc),
                    "desktop_context": exc.details,
                    "evidence_jsonl": str(evidence),
                }, ensure_ascii=False, indent=2), flush=True)
                return 3
        if args.build_test_ledger:
            evidence_entries = derive_ledger(ROOT / "runs")
            manual = load_manual_registrations(args.manual_registration)
            entries = merge_manual_registrations(evidence_entries, manual)
            validation_error = None
            try:
                validate_ledger(entries)
            except ValueError as exc:
                validation_error = str(exc)
            payload = {
                "schema_version": 2,
                "generated_at": datetime.now().astimezone().isoformat(
                    timespec="seconds"),
                "source": str((ROOT / "runs").resolve()),
                "manual_registration_source": str(
                    args.manual_registration.resolve()),
                "valid": validation_error is None,
                "validation_error": validation_error,
                "entries": entries,
            }
            args.ledger_output.parent.mkdir(parents=True, exist_ok=True)
            args.ledger_output.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8")
            print(json.dumps({
                "status": "OK" if validation_error is None else "INVALID",
                "entries": len(entries),
                "output": str(args.ledger_output.resolve()),
                "valid": validation_error is None,
                "validation_error": validation_error,
            }, ensure_ascii=False, indent=2), flush=True)
            if validation_error:
                return 2
            return 0
        if args.scaffold_manual_registration:
            evidence_entries = derive_ledger(ROOT / "runs")
            result = scaffold_manual_registrations(
                args.manual_registration, evidence_entries)
            print(json.dumps({"status": "OK", **result},
                             ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.build_campaign_report:
            if not args.campaign_id:
                raise ValueError("--build-campaign-report 必须提供 --campaign-id")
            manual = load_manual_registrations(args.manual_registration)
            report = build_campaign_report(
                ROOT / "runs", args.campaign_id,
                manual_registrations=manual)
            args.campaign_report_output.parent.mkdir(parents=True, exist_ok=True)
            args.campaign_report_output.write_text(
                json.dumps(report, ensure_ascii=False, indent=2),
                encoding="utf-8")
            print(json.dumps({
                "status": "OK" if report["campaign_valid"] else "INVALID",
                "output": str(args.campaign_report_output.resolve()),
                "data_quality": report["data_quality"],
                "red_lines": report["red_lines"],
            }, ensure_ascii=False, indent=2), flush=True)
            return 0 if report["campaign_valid"] else 2

        live_mode = any((
            args.preflight_live,
            args.dismiss_unexpected_item_dialog,
            args.discard_current_test_draft,
            args.execute_create_live,
            args.verify_live,
            args.fill_live,
            args.resume_fill_live,
            args.resume_open_item_dialog_live,
            args.resume_quantity_live,
            args.validate_filled_live,
            args.inspect_open_form,
            args.probe_correlation_field_live,
        ))
        if live_mode:
            from e10_desktop_guard import (DesktopGuardError,
                                           require_interactive_e10)
            from e10_query import get_digiwin_pids
            try:
                decision = require_interactive_e10(
                    get_digiwin_pids(),
                    allow_start_e10=args.allow_start_e10)
                _RUNTIME_METADATA["desktop_guard"] = {
                    "code": decision.code,
                    "message": decision.message,
                    "details": decision.details,
                }
            except DesktopGuardError as exc:
                _RUNTIME_METADATA["desktop_guard"] = {
                    "code": exc.failure_code,
                    "message": str(exc),
                    "details": exc.details,
                }
                evidence = _record_desktop_guard_terminal(
                    args.report_dir, exc)
                print(json.dumps({
                    "status": "BLOCKED",
                    "failure_code": exc.failure_code,
                    "message": str(exc),
                    "evidence_jsonl": str(evidence),
                }, ensure_ascii=False, indent=2), file=sys.stderr, flush=True)
                return 3

            if args.execute_create_live:
                live_risk = RiskLevel.COMMIT
            elif any((args.verify_live, args.preflight_live,
                      args.inspect_open_form)):
                live_risk = RiskLevel.READ_ONLY
            else:
                live_risk = RiskLevel.REVERSIBLE_WRITE
            from e10_mutex import MutexError
            from e10_query import acquire_single_instance
            try:
                _mutex = acquire_single_instance(
                    risk=live_risk,
                    allow_readonly_degradation=(
                        args.allow_readonly_lock_degradation))
                mutex_environment = _mutex.as_environment()
                _RUNTIME_METADATA.update(mutex_environment)
            except MutexError as exc:
                details = exc.as_dict()
                _RUNTIME_METADATA.update({
                    "lock_namespace": details.get("lock_namespace"),
                    "mutex_denied_reason": details.get(
                        "mutex_denied_reason"),
                })
                evidence = _record_mutex_guard_terminal(
                    args.report_dir, exc, live_risk)
                print(json.dumps({
                    "status": "BLOCKED",
                    "failure_code": exc.failure_code,
                    "message": str(exc),
                    "lock": details,
                    "evidence_jsonl": str(evidence),
                }, ensure_ascii=False, indent=2), file=sys.stderr, flush=True)
                return 3
        if args.preflight_live:
            print(json.dumps(live_preflight(), ensure_ascii=False, indent=2),
                  flush=True)
            return 0
        if args.probe_correlation_field_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--probe-correlation-field-live 仅允许 "
                    "--account-set FRKTEST --human-present")
            result = probe_correlation_field_live(args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.dismiss_unexpected_item_dialog:
            print(json.dumps(dismiss_unexpected_item_lookup(),
                             ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.discard_current_test_draft:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--discard-current-test-draft 仅允许 "
                    "--account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            result = discard_current_test_draft(record, args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.execute_create_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--execute-create-live 仅允许 "
                    "--account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            dataset_epoch = require_dataset_epoch(args.dataset_epoch)
            try:
                request_no = resolve_request_no(record.request_no, args.request_no)
            except ValueError as exc:
                gate_journal = _record_request_gate_terminal(
                    args.report_dir, status="FAILED", message=str(exc),
                    request_no=args.request_no or record.request_no,
                    payload_fingerprint=record.payload_fingerprint,
                    dataset_epoch=dataset_epoch)
                print(f"[防重闸] 证据: {gate_journal}", file=sys.stderr, flush=True)
                raise
            registry = RequestRegistry(args.request_registry)
            orphan_audit = audit_orphaned_runs(
                ROOT / "runs", request_no=request_no,
                dataset_epoch=dataset_epoch,
                registry_entries=registry.read())
            presave_orphans = [
                row for row in orphan_audit
                if row["action"] ==
                OrphanAction.PRE_SAVE_DRAFT_CHECK_REQUIRED.value]
            if presave_orphans:
                resolution = _resolve_presave_orphans(
                    record, args.report_dir, request_no=request_no,
                    dataset_epoch=dataset_epoch, orphans=presave_orphans)
                print(json.dumps({
                    "status": "ORPHAN_RESOLVED",
                    **resolution,
                }, ensure_ascii=False, indent=2), flush=True)
            decision = registry.decide(
                request_no, record.payload_fingerprint, args.campaign_id,
                retry_not_applied=args.retry_not_applied,
                dataset_epoch=dataset_epoch)
            if decision.action == GateAction.REJECT:
                gate_journal = _record_request_gate_terminal(
                    args.report_dir, status="FAILED",
                    message=f"业务请求防重闸拒绝: {decision.reason}",
                    request_no=request_no,
                    dataset_epoch=dataset_epoch,
                    payload_fingerprint=record.payload_fingerprint,
                    doc_no=(decision.existing or {}).get("doc_no"),
                    existing_evidence=(decision.existing or {}).get(
                        "evidence_jsonl"))
                print(f"[防重闸] 证据: {gate_journal}", file=sys.stderr, flush=True)
                raise RuntimeError(f"业务请求防重闸拒绝: {decision.reason}")
            if decision.action == GateAction.RETURN_CONFIRMED:
                gate_journal = _record_request_gate_terminal(
                    args.report_dir, status="CONFIRMED",
                    message="已有同 campaign、同请求号、同载荷确认记录；未接触 E10",
                    request_no=request_no,
                    dataset_epoch=dataset_epoch,
                    payload_fingerprint=record.payload_fingerprint,
                    doc_no=decision.existing.get("doc_no"),
                    existing_evidence=decision.existing.get("evidence_jsonl"))
                print(json.dumps({
                    "status": "SAVED_CONFIRMED",
                    "deduplicated": True,
                    "request_no": request_no,
                    "dataset_epoch": dataset_epoch,
                    "doc_no": decision.existing.get("doc_no"),
                    "evidence_jsonl": decision.existing.get("evidence_jsonl"),
                    "gate_journal": str(gate_journal),
                    "message": "已有同 campaign、同请求号、同载荷确认记录；未接触 E10",
                }, ensure_ascii=False, indent=2), flush=True)
                return 0
            if decision.action == GateAction.RECONCILE_ONLY:
                existing_doc = decision.existing.get("doc_no")
                if not existing_doc:
                    raise RuntimeError(
                        "既有 PENDING/UNKNOWN 无单号，只能人工核对；禁止再次保存")
                result = verify_requisition_live(
                    existing_doc, args.report_dir, request_no=request_no,
                    payload_fingerprint=record.payload_fingerprint,
                    dataset_epoch=dataset_epoch)
                registry.append(
                    request_no=request_no,
                    campaign_id=decision.campaign_id,
                    dataset_epoch=dataset_epoch,
                    payload_fingerprint=record.payload_fingerprint,
                    status=RequestStatus(result["status"]),
                    run_id=Path(result["journal"]).stem,
                    attempt=decision.attempt,
                    doc_no=existing_doc,
                    evidence_jsonl=str(Path(result["journal"]).resolve()),
                    details={"reconcile_only": True, "external": False},
                )
                result["request_gate"] = "RECONCILE_ONLY"
                result["request_no"] = request_no
                print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
                return 0 if result["status"] != "UNKNOWN" else 2
            result = execute_create_from_main(
                record, args.report_dir, input_path=args.input,
                request_no=request_no, dataset_epoch=dataset_epoch,
                request_registry=registry,
                gate_decision=decision)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.verify_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--verify-live 仅允许 --account-set FRKTEST --human-present")
            if not args.doc_no:
                raise RuntimeError("--verify-live 必须提供 --doc-no")
            if bool(args.request_no) != bool(args.dataset_epoch):
                raise RuntimeError(
                    "强身份核对必须同时提供 --request-no 与 --dataset-epoch")
            result = verify_requisition_live(
                args.doc_no, args.report_dir,
                request_no=args.request_no,
                dataset_epoch=args.dataset_epoch)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.fill_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--fill-live 仅允许 --account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            result = fill_live_form(record, args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.resume_fill_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--resume-fill-live 仅允许 --account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            result = resume_live_form(record, args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.resume_open_item_dialog_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--resume-open-item-dialog-live 仅允许 "
                    "--account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            result = resume_open_item_dialog(record, args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.resume_quantity_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--resume-quantity-live 仅允许 "
                    "--account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            result = resume_quantity(record, args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.validate_filled_live:
            if args.account_set != "FRKTEST" or not args.human_present:
                raise RuntimeError(
                    "--validate-filled-live 仅允许 "
                    "--account-set FRKTEST --human-present")
            record = load_one_record(args.input, args.sheet)
            result = validate_filled_form(record, args.report_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
            return 0
        if args.inspect_open_form:
            output = inspect_open_form(
                args.report_dir, timeout_seconds=args.probe_timeout)
            _progress(f"只读探针正常结束，证据文件：{output}")
            return 0

        record = load_one_record(args.input, args.sheet)
        result = {
            "input": str(args.input.resolve()),
            "record": record.as_business_mapping(),
            "request_no": record.request_no,
            "payload_fingerprint": record.payload_fingerprint,
        }
        if not args.validate_only:
            recording = args.recording or latest_recording(DEFAULT_RECORDINGS)
            result["recording_alignment"] = (
                align_recording(recording, record) if recording else None)
            result["semantic_plan"] = semantic_plan(record)
            result["execution_ready"] = record.warehouse is None
            result["blocked_by"] = (
                [] if record.warehouse is None
                else ["EXPLICIT_WAREHOUSE_SELECTION_UNIMPLEMENTED"])
            result["execution_constraints"] = [
                "FRKTEST_ONLY",
                "HUMAN_PRESENT_REQUIRED",
                "ONE_ROW_ONLY",
                "E10_UI_TEST_DATABASE_RECONCILE_EXTERNAL_FALSE",
                "NO_AUTOMATIC_RETRY_AFTER_SAVE",
            ]
        print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
        return 0
    except KeyboardInterrupt:
        print("\n[请购探针] 已由用户取消。", flush=True)
        return 130
    except Exception as exc:
        print(f"[请购探针] 失败：{type(exc).__name__}: {exc}",
              file=sys.stderr, flush=True)
        return 2
    finally:
        if _mutex is not None:
            _mutex.release()


if __name__ == "__main__":
    raise SystemExit(main())
